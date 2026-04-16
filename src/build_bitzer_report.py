#!/usr/bin/env python3
"""
build_bitzer_report.py
======================
Erstellt output/billing_report/bitzer_dinas_vergleich.xlsx
mit 4 Sheets: Zusammenfassung | PRE Dinas Detail | POST AX Detail | Cluster-Vergleich

Datenquellen:
  - output/bitzer_dinas_vergleich.xlsx   (PRE/POST already matched)
  - cluster_vergleich/406345_Bitzer*.xlsx (Soll EUR per Auftrag)
"""

import math, re
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC_XLSX    = Path('/home/user/TMS/output/bitzer_dinas_vergleich.xlsx')
CLUSTER_XLS = Path('/home/user/TMS/output/cluster_vergleich/406345_Bitzer_Kühlmaschinenbau_GmbH_cluster.xlsx')
REPORT_DIR  = Path('/home/user/TMS/output/billing_report')
OUT_XLSX    = REPORT_DIR / 'bitzer_dinas_vergleich.xlsx'

REPORT_DIR.mkdir(exist_ok=True)

# ── Gewichtsbänder ─────────────────────────────────────────────────────────────
BANDS = [50, 100, 150, 200, 250, 300, 500, 750, 1000, 2000, 3000]
def gew_band(kg):
    try:
        kg = float(kg)
    except (TypeError, ValueError):
        return '?'
    if math.isnan(kg) or kg <= 0:
        return '?'
    for b in BANDS:
        if kg <= b:
            return f'bis {b}kg'
    return 'über 3000kg'

def plz2(plz): return str(plz).strip()[:2]
def norm_id(v):
    try: return str(int(float(str(v))))
    except: return str(v).strip()

# ── Daten laden ────────────────────────────────────────────────────────────────
pre  = pd.read_excel(SRC_XLSX, sheet_name='PRE Dinas Detail',  header=2)
post = pd.read_excel(SRC_XLSX, sheet_name='POST AX Detail',    header=2)

for df in (pre, post):
    df['Tonnage (eff.)'] = pd.to_numeric(df['Tonnage (eff.)'], errors='coerce')
    df['Empfänger PLZ']  = df['Empfänger PLZ'].astype(str).str.strip()
    df['Empfänger Land'] = df['Empfänger Land'].astype(str).str.strip()
    df['_gwb']           = df['Tonnage (eff.)'].apply(gew_band)
    df['_plz2']          = df['Empfänger PLZ'].apply(plz2)
    df['_cluster']       = df['Empfänger Land'] + '|' + df['_plz2'] + '|' + df['_gwb']

# NK numerisch
PRE_NK  = ['Dinas Fracht','Dinas Diesel','Dinas Maut/SSD','Dinas Ausfuhr',
           'Dinas Verzollung','Dinas Zoll Duty','Dinas Zollbetrag',
           'Dinas Sulphur','Dinas Nebenkostenpausch.','Dinas Redebit',
           'Dinas Sonstige','Dinas Gesamt']
POST_NK = ['AX Fracht','AX Diesel','AX Maut','AX Nebengebühr',
           'AX Lademittel','AX Peak','AX EUST/Zoll','AX Versicherung','AX Gesamt']

for c in PRE_NK:
    pre[c] = pd.to_numeric(pre.get(c), errors='coerce')
for c in POST_NK:
    post[c] = pd.to_numeric(post.get(c), errors='coerce')

# ── Soll EUR aus Cluster-Excel joinen ─────────────────────────────────────────
cl_df = pd.read_excel(CLUSTER_XLS, sheet_name=0)
alt_cl = cl_df[cl_df['System'] == 'alt'][['Auftrags-Nr','Soll EUR']].copy()
neu_cl = cl_df[cl_df['System'] == 'neu'][['Auftrags-Nr','Soll EUR']].copy()
alt_cl['_aid'] = alt_cl['Auftrags-Nr'].apply(norm_id)
neu_cl['_aid'] = neu_cl['Auftrags-Nr'].apply(norm_id)
alt_cl['Soll EUR'] = pd.to_numeric(alt_cl['Soll EUR'], errors='coerce')
neu_cl['Soll EUR'] = pd.to_numeric(neu_cl['Soll EUR'], errors='coerce')
soll_pre  = alt_cl.dropna(subset=['Soll EUR']).set_index('_aid')['Soll EUR'].to_dict()
soll_post = neu_cl.dropna(subset=['Soll EUR']).set_index('_aid')['Soll EUR'].to_dict()
pre['_aid']    = pre['Auftragsnummer'].apply(norm_id)
post['_aid']   = post['Auftragsnummer'].apply(norm_id)
pre['Soll EUR']  = pre['_aid'].map(soll_pre)
post['Soll EUR'] = post['_aid'].map(soll_post)
print(f'Soll EUR matched: PRE {pre["Soll EUR"].notna().sum()}/{len(pre)}, '
      f'POST {post["Soll EUR"].notna().sum()}/{len(post)}')

# ── Cluster-Statistiken ────────────────────────────────────────────────────────
def cluster_stats(pre, post):
    pa = pre.groupby('_cluster').agg(
        n_pre=('Dinas Gesamt','count'),
        avg_d_ges=('Dinas Gesamt','mean'),
        avg_d_fr=('Dinas Fracht','mean'),
        avg_d_di=('Dinas Diesel','mean'),
        avg_d_maut=('Dinas Maut/SSD','mean'),
    ).reset_index()
    oa = post.groupby('_cluster').agg(
        n_post=('AX Gesamt','count'),
        avg_ax_ges=('AX Gesamt','mean'),
        avg_ax_fr=('AX Fracht','mean'),
        avg_ax_di=('AX Diesel','mean'),
        avg_ax_maut=('AX Maut','mean'),
        avg_ax_neben=('AX Nebengebühr','mean'),
    ).reset_index()
    s = pa.merge(oa, on='_cluster', how='inner')
    s = s[s['avg_ax_ges'] < s['avg_d_ges']].copy()
    s['delta']     = s['avg_ax_ges'] - s['avg_d_ges']
    s['delta_pct'] = s['delta'] / s['avg_d_ges'].replace(0, np.nan) * 100
    s['delta_fr']  = s['avg_ax_fr']   - s['avg_d_fr']
    s['delta_di']  = s['avg_ax_di']   - s['avg_d_di']
    s['delta_maut']= s['avg_ax_maut'] - s['avg_d_maut']
    s['est_loss']  = s['delta'] * s['n_post']
    def abw(r):
        neg = {k:v for k,v in {'Fracht':r.delta_fr,'Diesel':r.delta_di,
                                'Maut':r.delta_maut}.items() if v < -0.5}
        if not neg: return 'n/a'
        t = sum(neg.values())
        return 'Hauptursache: ' + ', '.join(
            f'{k}: {v/t*100:+.0f}%' for k,v in sorted(neg.items(), key=lambda x:x[1])[:2])
    s['abw_grund'] = s.apply(abw, axis=1)
    parts = s['_cluster'].str.split('|', expand=True)
    s['Land'] = parts[0]; s['PLZ_pre'] = parts[1]; s['Gew_band'] = parts[2]
    return s.sort_values('est_loss').reset_index(drop=True)

stats = cluster_stats(pre, post)
print(f'Unterfakturierungs-Cluster: {len(stats)}, '
      f'Σ Verlust: {stats["est_loss"].sum():,.0f} EUR')

# ── Kontroll-Sendung ───────────────────────────────────────────────────────────
def find_kontroll(pre_all, post_all):
    if pre_all.empty or post_all.empty: return None, None
    best = (None, None, float('inf'))
    for pi, pr in pre_all.iterrows():
        pr_n = str(pr.get('Empfänger Name',''))[:12]
        pr_kg = pr.get('Tonnage (eff.)')
        if pd.isna(pr_kg): continue
        same = post_all[post_all['Empfänger Name'].astype(str).str[:12] == pr_n]
        for oi, po in (same if not same.empty else post_all).iterrows():
            po_kg = po.get('Tonnage (eff.)')
            if pd.isna(po_kg): continue
            d = abs(float(pr_kg) - float(po_kg))
            if d < best[2]: best = (pi, oi, d)
    return best[0], best[1]

# ── Styles ─────────────────────────────────────────────────────────────────────
def fill(h): return PatternFill('solid', fgColor=h)
def font(bold=False, color='000000', size=9, italic=False):
    return Font(bold=bold, color=color, size=size, italic=italic)
THIN = Side(border_style='thin', color='AAAAAA')
MED  = Side(border_style='medium', color='555555')
BRD  = Border(left=THIN, right=THIN, top=THIN,  bottom=THIN)
BRD_H= Border(left=THIN, right=THIN, top=MED,   bottom=MED)
EUR  = '#,##0.00'

def wc(ws, r, c, v=None, f=None, fn=None, al='left', fmt=None, brd=None):
    cell = ws.cell(row=r, column=c)
    if v is not None: cell.value = v
    if f:  cell.fill  = f
    if fn: cell.font  = fn
    if fmt: cell.number_format = fmt
    if brd: cell.border = brd
    cell.alignment = Alignment(horizontal=al, vertical='center')
    return cell

def mrow(ws, r, c1, c2, v, f=None, fn=None, al='left'):
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    wc(ws, r, c1, v, f, fn, al)

# ── Spalten-Layout ─────────────────────────────────────────────────────────────
COLS = [
    ('RN',       'Rechnungsnummer','Rechnungsnummer', 11),
    ('Auftrag',  'Auftragsnummer', 'Auftragsnummer',  12),
    ('Datum',    'Leistungsdatum', 'Leistungsdatum',  11),
    ('Empfänger','Empfänger Name', 'Empfänger Name',  22),
    ('PLZ',      'Empfänger PLZ',  'Empfänger PLZ',    7),
    ('Land',     'Empfänger Land', 'Empfänger Land',   5),
    ('Ton. kg',  'Tonnage (eff.)', 'Tonnage (eff.)',   9),
    ('Soll €',   'Soll EUR',       'Soll EUR',         9),
    # NK (shared physical columns, different labels for PRE vs POST)
    ('Fracht €',    'Dinas Fracht',          'AX Fracht',        10),
    ('Diesel €',    'Dinas Diesel',          'AX Diesel',         9),
    ('Maut €',      'Dinas Maut/SSD',        'AX Maut',           9),
    ('Ausfuhr €',   'Dinas Ausfuhr',         'AX Nebengebühr',    9),
    ('Verzoll. €',  'Dinas Verzollung',      'AX Lademittel',     9),
    ('ZollD. €',    'Dinas Zoll Duty',       'AX Peak',           8),
    ('Sulphur €',   'Dinas Sulphur',         'AX EUST/Zoll',      9),
    ('NKP. €',      'Dinas Nebenkostenpausch.','AX Versicherung', 9),
    ('Redebit €',   'Dinas Redebit',         None,                8),
    ('Sonstige €',  'Dinas Sonstige',        None,                8),
    ('GESAMT €',    'Dinas Gesamt',          'AX Gesamt',        11),
]
N = len(COLS)

PRE_OVR  = {9:'D-Fracht €',10:'D-Diesel €',11:'D-Maut €',12:'D-Ausfuhr €',
            13:'D-Verzoll €',14:'D-ZollD €',15:'D-Sulphur €',16:'D-NKP €',
            17:'D-Redebit €',18:'D-Sonstige €',19:'D-GESAMT €'}
POST_OVR = {9:'AX Fracht €',10:'AX Diesel €',11:'AX Maut €',
            12:'AX Neben €',13:'AX Lademl €',14:'AX Peak €',
            15:'AX EUST €',16:'AX Versich €',17:'',18:'',19:'AX GESAMT €'}

def write_hdrs(ws, row, ovr, hf):
    for ci,(h,_,_,_) in enumerate(COLS, 1):
        wc(ws, row, ci, ovr.get(ci,h), hf,
           font(bold=True,color='FFFFFF',size=8),'center',brd=BRD_H)

def write_drow(ws, row, src, si, rf):
    for ci,cd in enumerate(COLS, 1):
        col = cd[si]
        v   = src.get(col) if col else None
        if v is not None and pd.isna(v): v = None
        e = '€' in cd[0]
        wc(ws, row, ci, v, rf, font(size=9),
           'right' if e or cd[0]=='Ton. kg' else 'left',
           EUR if e else None, BRD)

# ── Sheet 1: Zusammenfassung ───────────────────────────────────────────────────
def ws_zusammenfassung(ws):
    ws.title = 'Zusammenfassung'
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 22
    rows = [
        ('Bitzer Kühlmaschinenbau GmbH (406345) — Dinas PDF vs AX Vergleich', ''),
        ('', ''),
        ('── PRE (Dinas-System) ──────────────────────────', ''),
        ('BI-Zeilen PRE (gematcht)', len(pre)),
        ('Σ Dinas Fracht EUR', round(pre['Dinas Fracht'].sum(), 2)),
        ('Σ Dinas Gesamt EUR', round(pre['Dinas Gesamt'].sum(), 2)),
        ('', ''),
        ('── POST (AX-System) ────────────────────────────', ''),
        ('BI-Zeilen POST', len(post)),
        ('Σ AX Fracht EUR', round(post['AX Fracht'].sum(), 2)),
        ('Σ AX Gesamt EUR', round(post['AX Gesamt'].sum(), 2)),
        ('', ''),
        ('── Cluster-Vergleich ────────────────────────────', ''),
        ('Unterfakturierungs-Cluster', len(stats)),
        ('Σ geschätzter Gesamtverlust EUR', round(stats['est_loss'].sum(), 0)),
    ]
    for i,(k,v) in enumerate(rows, 1):
        ws.cell(i,1).value = k
        ws.cell(i,2).value = v
        if v=='' and k.startswith('─'):
            ws.cell(i,1).fill = fill('D6E4F7')
            ws.cell(i,1).font = font(bold=True)
        elif i==1:
            ws.cell(i,1).font = font(bold=True,size=12)
        else:
            ws.cell(i,2).alignment = Alignment(horizontal='right')

# ── Sheet 4: Cluster-Vergleich ─────────────────────────────────────────────────
def ws_cluster(ws, top_n=20):
    ws.title = 'Cluster-Vergleich'
    mrow(ws,1,1,N,'Bitzer Kühlmaschinenbau GmbH (406345) — Cluster-Vergleich: Unterfakturierung',
         fill('1F497D'), font(bold=True,color='FFFFFF',size=12),'center')
    mrow(ws,2,1,N,'PRE: Dinas NK  |  POST: AX NK  |  Soll EUR aus Tarifmotor  |  Grün = Kontroll-Paar',
         fill('2E4057'), font(color='CCCCCC',size=9,italic=True),'center')
    row = 2

    for _,cl in stats.head(top_n).iterrows():
        ckey = cl['_cluster']
        pre_all  = pre[pre['_cluster']==ckey]
        post_all = post[post['_cluster']==ckey]
        pre_s    = pre_all.head(5)
        post_s   = post_all.head(5)
        kpi, koi = find_kontroll(pre_all, post_all)

        row += 1
        mrow(ws,row,1,N,
             f'  CLUSTER: Land={cl.Land}  |  PLZ={cl.PLZ_pre}  |  '
             f'Gew.band={cl.Gew_band}  |  n PRE={int(cl.n_pre)}  |  n POST={int(cl.n_post)}',
             fill('2E75B6'), font(bold=True,color='FFFFFF',size=10))
        row += 1
        mrow(ws,row,1,N,
             f'  Ø Dinas: {cl.avg_d_ges:,.2f} €  |  Ø AX: {cl.avg_ax_ges:,.2f} €  '
             f'|  Δ: {cl.delta:,.2f} € ({cl.delta_pct:+.1f}%)  '
             f'|  Gesch. Verlust: {cl.est_loss:,.0f} €  |  {cl.abw_grund}',
             fill('D6E4F7'), font(size=9))

        # PRE
        row += 1
        mrow(ws,row,1,N,'  PRE — Dinas Rechnungen',fill('4472C4'),font(bold=True,color='FFFFFF'))
        row += 1; write_hdrs(ws, row, PRE_OVR, fill('4472C4'))
        for _,pr in pre_s.iterrows():
            row += 1; write_drow(ws, row, pr, 1, fill('BDD7EE'))

        # POST
        row += 1
        mrow(ws,row,1,N,'  POST — AX Rechnungen',fill('C55A11'),font(bold=True,color='FFFFFF'))
        row += 1; write_hdrs(ws, row, POST_OVR, fill('C55A11'))
        for _,po in post_s.iterrows():
            row += 1; write_drow(ws, row, po, 2, fill('FCE4D6'))

        # Kontroll
        if kpi is not None or koi is not None:
            row += 1
            mrow(ws,row,1,N,'  ★ Kontroll-Sendung (bestes PRE/POST-Paar)',
                 fill('E2EFDA'), font(bold=True,color='1F5C00'))
            if kpi is not None:
                row += 1
                wc(ws,row,1,'PRE',fill('E2EFDA'),font(bold=True,italic=True))
                pr = pre.loc[kpi]
                for ci,(h,ps,_,_) in enumerate(COLS[1:],2):
                    v=pr.get(ps) if ps else None
                    if v is not None and pd.isna(v): v=None
                    e='€' in h
                    wc(ws,row,ci,v,fill('E2EFDA'),font(size=9),
                       'right' if e else 'left', EUR if e else None, BRD)
            if koi is not None:
                row += 1
                wc(ws,row,1,'POST',fill('E2EFDA'),font(bold=True,italic=True))
                po = post.loc[koi]
                for ci,(h,_,qs,_) in enumerate(COLS[1:],2):
                    v=po.get(qs) if qs else None
                    if v is not None and pd.isna(v): v=None
                    e='€' in h
                    wc(ws,row,ci,v,fill('E2EFDA'),font(size=9),
                       'right' if e else 'left', EUR if e else None, BRD)
        row += 1  # Leerzeile

    for ci,(_,_,_,w) in enumerate(COLS,1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = 'A3'
    return row

# ── Schreiben ──────────────────────────────────────────────────────────────────
wb = Workbook()
ws_zusammenfassung(wb.active)
ws_pre = wb.create_sheet('PRE Dinas Detail')
ws_pre.append(list(pre.columns))
for _,r in pre.iterrows(): ws_pre.append(list(r))
ws_post = wb.create_sheet('POST AX Detail')
ws_post.append(list(post.columns))
for _,r in post.iterrows(): ws_post.append(list(r))
last_row = ws_cluster(wb.create_sheet('Cluster-Vergleich'))

wb.save(OUT_XLSX)
print(f'\nGespeichert: {OUT_XLSX}')
print(f'  Cluster-Vergleich: {last_row} Zeilen')
print(f'  PRE Dinas Detail:  {len(pre)} Zeilen')
print(f'  POST AX Detail:    {len(post)} Zeilen')
