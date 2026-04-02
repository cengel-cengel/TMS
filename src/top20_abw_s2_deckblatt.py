"""
Script 2: Excel - Deckblatt + Abweichungsübersicht sheet (20 top deviations).
"""
import pickle, pandas as pd, numpy as np
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import (PatternFill, Font, Alignment, Border, Side,
                              DEFAULT_FONT)
from openpyxl.utils import get_column_letter

OUT = Path('/home/user/TMS/output')

# ── Load stats ────────────────────────────────────────────────────────────────
with open(OUT / 'top20_abw_stats.pkl', 'rb') as f:
    d = pickle.load(f)
stats = d['stats']
top20_dev = d['top20_deviations']
top20_meta = d['top20_meta']

# ── Styles ────────────────────────────────────────────────────────────────────
WHITE_FILL  = PatternFill(fill_type=None)  # no fill = white
BLUE_FILL   = PatternFill('solid', fgColor='1F4E79')
LBLUE_FILL  = PatternFill('solid', fgColor='2E75B6')
HEADER_FILL = PatternFill('solid', fgColor='D6E4F0')
RED_FILL    = PatternFill('solid', fgColor='FF0000')
ORANGE_FILL = PatternFill('solid', fgColor='FFC000')
YELLOW_FILL = PatternFill('solid', fgColor='FFFF00')
GREEN_FILL  = PatternFill('solid', fgColor='92D050')
GREY_FILL   = PatternFill('solid', fgColor='BFBFBF')

def hfont(bold=False, size=11, color='000000', name='Calibri'):
    return Font(name=name, bold=bold, size=size, color=color)

def thin_border():
    s = Side(style='thin')
    return Border(left=s, right=s, top=s, bottom=s)

def severity_fill(delta_pct, n_post):
    if n_post == 0:
        return GREY_FILL
    if pd.isna(delta_pct):
        return WHITE_FILL
    if delta_pct <= -40:
        return RED_FILL
    if delta_pct <= -20:
        return ORANGE_FILL
    if delta_pct <= -5:
        return YELLOW_FILL
    if delta_pct >= 10:
        return GREEN_FILL
    return WHITE_FILL

def set_col_width(ws, col, width):
    ws.column_dimensions[get_column_letter(col)].width = width

# ── Workbook ──────────────────────────────────────────────────────────────────
wb = Workbook()

# ═══════════════════════════════════════════════════════════════════════════════
# Sheet 1: Deckblatt
# ═══════════════════════════════════════════════════════════════════════════════
ws_deck = wb.active
ws_deck.title = 'Deckblatt'

ws_deck.column_dimensions['A'].width = 28
ws_deck.column_dimensions['B'].width = 60

row = 1
ws_deck.row_dimensions[row].height = 40
c = ws_deck.cell(row, 1, 'TOP-20 ABWEICHUNGSANALYSE')
c.font = hfont(bold=True, size=18, color='FFFFFF')
c.fill = BLUE_FILL
c.alignment = Alignment(horizontal='center', vertical='center')
ws_deck.merge_cells(f'A{row}:B{row}')

row += 1
ws_deck.cell(row, 1, 'TMS Migration: Dinas → Microsoft AX').font = hfont(bold=True, size=13)
ws_deck.merge_cells(f'A{row}:B{row}')

row += 1
ws_deck.cell(row, 1, 'Erstellungsdatum:').font = hfont(bold=True)
ws_deck.cell(row, 2, '02.04.2026')

row += 1
ws_deck.cell(row, 1, 'Migration Go-Live:').font = hfont(bold=True)
ws_deck.cell(row, 2, '26.09.2025')

row += 1
ws_deck.cell(row, 1, 'Analysezeitraum PRE:').font = hfont(bold=True)
ws_deck.cell(row, 2, 'vor 26.09.2025')

row += 1
ws_deck.cell(row, 1, 'Analysezeitraum POST:').font = hfont(bold=True)
ws_deck.cell(row, 2, 'ab 26.09.2025')

row += 1
ws_deck.cell(row, 1, 'Datengrundlage:').font = hfont(bold=True)
ws_deck.cell(row, 2, 'Nur Sendungen mit Rechnungsnummer (fakturiert)')

row += 1
ws_deck.cell(row, 1, 'Kundenbasis:').font = hfont(bold=True)
ws_deck.cell(row, 2, f'Top {len(top20_meta)} Kunden nach Gesamterlös')

row += 2
ws_deck.cell(row, 1, 'Methodik').font = hfont(bold=True, size=12)
ws_deck.merge_cells(f'A{row}:B{row}')

row += 1
ws_deck.cell(row, 1, '').fill = WHITE_FILL
ws_deck.cell(row, 2, (
    'Für jeden Kunden × Verkehrsart werden Durchschnittserlöse je Sendung '
    'im PRE- und POST-Zeitraum verglichen. Die Abweichung (Δ€, Δ%) gibt an, '
    'wie stark sich der durchschnittliche Erlös pro Sendung verändert hat. '
    'Severity = |Δ€| × max(n_pre, n_post) gewichtet nach Sendungsvolumen.'
))
ws_deck.cell(row, 2).alignment = Alignment(wrap_text=True)
ws_deck.row_dimensions[row].height = 45
ws_deck.merge_cells(f'B{row}:B{row}')

row += 2
ws_deck.cell(row, 1, 'VERKEHRSART MAPPING').font = hfont(bold=True, size=12)
ws_deck.merge_cells(f'A{row}:B{row}')

row += 1
hdr_row = row
ws_deck.cell(row, 1, 'va_norm').font = hfont(bold=True, color='FFFFFF')
ws_deck.cell(row, 1).fill = LBLUE_FILL
ws_deck.cell(row, 2, 'Dinas (PRE)  →  AX (POST)').font = hfont(bold=True, color='FFFFFF')
ws_deck.cell(row, 2).fill = LBLUE_FILL

mapping = [
    ('SA',             'SA (national)  →  SA (national)'),
    ('Export',         'Export/Charter International + Stgt  →  Export'),
    ('Charter',        'Charter International + TP/DP  →  Charter International + Charter'),
    ('Charter National','Charter National  →  Charter National'),
    ('Belog SE',       'Belog + Belog SE  →  Belog'),
    ('SE',             'SE  →  SE'),
]
for va, desc in mapping:
    row += 1
    ws_deck.cell(row, 1, va)
    ws_deck.cell(row, 2, desc)

row += 2
ws_deck.cell(row, 1, 'FARBKODIERUNG ABWEICHUNGEN').font = hfont(bold=True, size=12)
ws_deck.merge_cells(f'A{row}:B{row}')

legend = [
    (RED_FILL,    'Δ% ≤ -40%  –  Kritische Abweichung'),
    (ORANGE_FILL, 'Δ% -40% bis -20%  –  Starke Abweichung'),
    (YELLOW_FILL, 'Δ% -20% bis -5%  –  Moderate Abweichung'),
    (GREEN_FILL,  'Δ% ≥ +10%  –  Erlössteigerung'),
    (GREY_FILL,   'n_post = 0  –  Verkehr / Kunde verloren'),
]
for fill, label in legend:
    row += 1
    c = ws_deck.cell(row, 1, '   ')
    c.fill = fill
    ws_deck.cell(row, 2, label)

# ═══════════════════════════════════════════════════════════════════════════════
# Sheet 2: Abweichungsübersicht (top 20)
# ═══════════════════════════════════════════════════════════════════════════════
ws_ov = wb.create_sheet('Abweichungsübersicht')

headers = [
    'Rang', 'Kunden Nr', 'Kunden Name', 'Verkehrsart (norm.)',
    'n PRE', 'n POST',
    'Ø Erlös PRE (€)', 'Ø Erlös POST (€)', 'Δ € / Sendung', 'Δ %',
    'Ø DB II PRE (€)', 'Ø DB II POST (€)',
    'Severity Score', 'Status'
]
col_widths = [6, 12, 36, 22, 8, 8, 18, 18, 16, 10, 18, 18, 16, 22]

for i, (h, w) in enumerate(zip(headers, col_widths), 1):
    c = ws_ov.cell(1, i, h)
    c.font = hfont(bold=True, color='FFFFFF')
    c.fill = LBLUE_FILL
    c.alignment = Alignment(horizontal='center', wrap_text=True)
    c.border = thin_border()
    set_col_width(ws_ov, i, w)

ws_ov.row_dimensions[1].height = 30
ws_ov.freeze_panes = 'A2'

# Also get totals for lost-customer detection
lost_customers = set()
tot = stats[stats['va_norm'] == '(Gesamt)']
for _, r in tot.iterrows():
    if r['n_post'] == 0:
        lost_customers.add(r['Kunden Nr BK'])

for rank, (_, row_data) in enumerate(top20_dev.iterrows(), 1):
    r = rank + 1
    n_pre  = int(row_data['n_pre'])  if not pd.isna(row_data['n_pre'])  else 0
    n_post = int(row_data['n_post']) if not pd.isna(row_data['n_post']) else 0
    avg_pre  = row_data['avg_pre']
    avg_post = row_data['avg_post']
    delta_eur = row_data['delta_eur']
    delta_pct = row_data['delta_pct']
    sev = row_data['severity']

    if n_post == 0:
        status = 'VERKEHR VERLOREN'
    elif pd.isna(delta_pct):
        status = 'Neu (kein PRE)'
    elif delta_pct <= -40:
        status = 'Kritisch'
    elif delta_pct <= -20:
        status = 'Stark'
    elif delta_pct <= -5:
        status = 'Moderat'
    elif delta_pct >= 10:
        status = 'Steigerung'
    else:
        status = 'Stabil'

    fill = severity_fill(delta_pct, n_post)

    def fmt_eur(v):
        return round(float(v), 2) if not pd.isna(v) else ''
    def fmt_pct(v):
        return round(float(v), 2) if not pd.isna(v) else ''

    vals = [
        rank,
        row_data['Kunden Nr BK'],
        row_data['Kunden Name'],
        row_data['va_norm'],
        n_pre, n_post,
        fmt_eur(avg_pre),
        fmt_eur(avg_post),
        fmt_eur(delta_eur),
        fmt_pct(delta_pct),
        fmt_eur(row_data.get('avg_pre_dbii', np.nan)),
        fmt_eur(row_data.get('avg_post_dbii', np.nan)),
        round(float(sev), 0) if not pd.isna(sev) else '',
        status
    ]

    num_formats = {
        7: '#,##0.00 "€"',
        8: '#,##0.00 "€"',
        9: '#,##0.00 "€"',
        10: '0.00"%"',
        11: '#,##0.00 "€"',
        12: '#,##0.00 "€"',
        13: '#,##0',
    }

    for col_i, val in enumerate(vals, 1):
        c = ws_ov.cell(r, col_i, val)
        c.fill = fill
        c.border = thin_border()
        c.alignment = Alignment(horizontal='center')
        if col_i in num_formats and isinstance(val, (int, float)) and val != '':
            c.number_format = num_formats[col_i]

print("Deckblatt and Abweichungsübersicht written")

# ── Save intermediate workbook ─────────────────────────────────────────────────
xlsx_path = OUT / 'top20_audit_abweichungsanalyse_2026.xlsx'
wb.save(str(xlsx_path))
print(f"Saved: {xlsx_path}")
