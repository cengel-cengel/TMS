"""
Script 3: Add per-customer detail sheets to the Excel workbook.
One sheet per top-20 customer: summary KPI by va_norm + 5 PRE + 5 POST samples.
"""
import pickle, pandas as pd, numpy as np
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = Path('/home/user/TMS/output')
XLSX = OUT / 'top20_audit_abweichungsanalyse_2026.xlsx'

# ── Load ──────────────────────────────────────────────────────────────────────
with open(OUT / 'top20_abw_stats.pkl', 'rb') as f:
    d = pickle.load(f)
stats      = d['stats']
top20_meta = d['top20_meta']

with open(OUT / 'top20_abw_samples.pkl', 'rb') as f:
    samples_df = pickle.load(f)

wb = load_workbook(str(XLSX))

# ── Styles ────────────────────────────────────────────────────────────────────
WHITE_FILL  = PatternFill(fill_type=None)
BLUE_FILL   = PatternFill('solid', fgColor='1F4E79')
LBLUE_FILL  = PatternFill('solid', fgColor='2E75B6')
PRE_FILL    = PatternFill('solid', fgColor='D6E4F0')
POST_FILL   = PatternFill('solid', fgColor='E2EFDA')
GREY_FILL   = PatternFill('solid', fgColor='BFBFBF')
RED_FILL    = PatternFill('solid', fgColor='FF0000')
ORANGE_FILL = PatternFill('solid', fgColor='FFC000')
YELLOW_FILL = PatternFill('solid', fgColor='FFFF00')

def hfont(bold=False, size=11, color='000000'):
    return Font(name='Calibri', bold=bold, size=size, color=color)

def thin():
    s = Side(style='thin')
    return Border(left=s, right=s, top=s, bottom=s)

def sev_fill(delta_pct, n_post):
    if n_post == 0:            return GREY_FILL
    if pd.isna(delta_pct):     return WHITE_FILL
    if delta_pct <= -40:       return RED_FILL
    if delta_pct <= -20:       return ORANGE_FILL
    if delta_pct <= -5:        return YELLOW_FILL
    return WHITE_FILL

def set_width(ws, col, w):
    ws.column_dimensions[get_column_letter(col)].width = w

# ── Sample columns to display ─────────────────────────────────────────────────
SHOW_COLS = [
    'Leistungsdatum', 'Rechnungsnummer', 'Ausgangsbordero',
    'va_norm', 'Verkehrsart', 'Verkehrsart Detail',
    'Versender PLZ', 'Empfänger Land', 'Empfänger PLZ', 'route_2',
    'Colli', 'Tonnage (eff.)', 'Lademeter', 'Volumen',
    'Erlöse Fracht', 'Erlöse Diesel', 'Erlöse Maut', 'Erlöse Nebengebühr',
    'Erloese', 'Kosten Gesamt', 'DB II', 'DB II%',
]
SHOW_COLS = [c for c in SHOW_COLS if c in samples_df.columns]

NUM_COLS = {'Colli','Tonnage (eff.)','Lademeter','Volumen',
            'Erlöse Fracht','Erlöse Diesel','Erlöse Maut','Erlöse Nebengebühr',
            'Erloese','Kosten Gesamt','DB II','DB II%'}

COL_WIDTHS = {
    'Leistungsdatum': 14, 'Rechnungsnummer': 16, 'Ausgangsbordero': 18,
    'va_norm': 14, 'Verkehrsart': 20, 'Verkehrsart Detail': 16,
    'Versender PLZ': 12, 'Empfänger Land': 14, 'Empfänger PLZ': 12,
    'route_2': 14, 'Colli': 7, 'Tonnage (eff.)': 14, 'Lademeter': 11,
    'Volumen': 10, 'Erlöse Fracht': 15, 'Erlöse Diesel': 14,
    'Erlöse Maut': 13, 'Erlöse Nebengebühr': 18,
    'Erloese': 13, 'Kosten Gesamt': 14, 'DB II': 12, 'DB II%': 10,
}

SUMMARY_HEADERS = [
    'Verkehrsart', 'n PRE', 'n POST',
    'Ø Erlös PRE €', 'Ø Erlös POST €', 'Δ € / Sendung', 'Δ %',
    'Ø DB II PRE €', 'Ø DB II POST €', 'Status'
]

# ── Build one sheet per customer ──────────────────────────────────────────────
for _, meta_row in top20_meta.iterrows():
    knr   = meta_row['kunden_nr']
    kname = str(meta_row['kunden_name'])

    # Sheet name max 31 chars
    sname = kname[:28].strip().rstrip('.')
    if sname in wb.sheetnames:
        sname = sname[:25] + f'_{knr}'[:5]
    ws = wb.create_sheet(sname)

    # ── Title row ──────────────────────────────────────────────────────────
    ws.merge_cells('A1:V1')
    c = ws.cell(1, 1, f'{kname}  |  Kunden-Nr: {knr}')
    c.font = hfont(bold=True, size=14, color='FFFFFF')
    c.fill = BLUE_FILL
    c.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 28

    # ── Total Erlös from meta ──────────────────────────────────────────────
    gesamt_erloes = meta_row.get('gesamt_erloes', np.nan)
    ws.cell(2, 1, 'Gesamterlös (POST):').font = hfont(bold=True)
    c2 = ws.cell(2, 2, round(float(gesamt_erloes), 2) if not pd.isna(gesamt_erloes) else '')
    if isinstance(c2.value, float): c2.number_format = '#,##0.00 "€"'

    # ── Section A: Summary by va_norm ──────────────────────────────────────
    row = 4
    ws.merge_cells(f'A{row}:J{row}')
    c = ws.cell(row, 1, 'A — Kennzahlen nach Verkehrsart (PRE vs. POST)')
    c.font = hfont(bold=True, size=12, color='FFFFFF')
    c.fill = LBLUE_FILL

    row += 1
    for ci, h in enumerate(SUMMARY_HEADERS, 1):
        c = ws.cell(row, ci, h)
        c.font = hfont(bold=True, color='FFFFFF')
        c.fill = LBLUE_FILL
        c.border = thin()
        c.alignment = Alignment(horizontal='center', wrap_text=True)
    ws.row_dimensions[row].height = 28

    cust_stats = stats[(stats['Kunden Nr BK'] == knr) & (stats['va_norm'] != '(Gesamt)')]
    has_any_post = False

    for _, sr in cust_stats.iterrows():
        row += 1
        n_pre  = int(sr['n_pre'])  if not pd.isna(sr['n_pre'])  else 0
        n_post = int(sr['n_post']) if not pd.isna(sr['n_post']) else 0
        if n_post > 0:
            has_any_post = True

        delta_pct = sr['delta_pct']
        fill = sev_fill(delta_pct, n_post)

        status = ('VERKEHR VERLOREN' if n_post == 0
                  else 'Kritisch'  if not pd.isna(delta_pct) and delta_pct <= -40
                  else 'Stark'     if not pd.isna(delta_pct) and delta_pct <= -20
                  else 'Moderat'   if not pd.isna(delta_pct) and delta_pct <= -5
                  else 'Stabil')

        def fe(v): return round(float(v), 2) if not pd.isna(v) else ''
        def fp(v): return round(float(v), 2) if not pd.isna(v) else ''

        row_vals = [
            sr['va_norm'], n_pre, n_post,
            fe(sr['avg_pre']), fe(sr['avg_post']), fe(sr['delta_eur']), fp(delta_pct),
            fe(sr.get('avg_pre_dbii', np.nan)), fe(sr.get('avg_post_dbii', np.nan)),
            status
        ]
        num_fmt_map = {4:'#,##0.00 "€"', 5:'#,##0.00 "€"', 6:'#,##0.00 "€"',
                       7:'0.00"%"', 8:'#,##0.00 "€"', 9:'#,##0.00 "€"'}
        for ci, val in enumerate(row_vals, 1):
            c = ws.cell(row, ci, val)
            c.fill = fill
            c.border = thin()
            c.alignment = Alignment(horizontal='center')
            if ci in num_fmt_map and isinstance(val, (int, float)) and val != '':
                c.number_format = num_fmt_map[ci]

    # Totals row
    tot_row = stats[(stats['Kunden Nr BK'] == knr) & (stats['va_norm'] == '(Gesamt)')]
    if len(tot_row) > 0:
        row += 1
        tr = tot_row.iloc[0]
        n_pre_t  = int(tr['n_pre'])  if not pd.isna(tr['n_pre'])  else 0
        n_post_t = int(tr['n_post']) if not pd.isna(tr['n_post']) else 0
        fill_t = sev_fill(tr['delta_pct'], n_post_t)
        tot_vals = [
            'GESAMT', n_pre_t, n_post_t,
            fe(tr['avg_pre']), fe(tr['avg_post']), fe(tr['delta_eur']), fp(tr['delta_pct']),
            fe(tr.get('avg_pre_dbii', np.nan)), fe(tr.get('avg_post_dbii', np.nan)), ''
        ]
        for ci, val in enumerate(tot_vals, 1):
            c = ws.cell(row, ci, val)
            c.fill = fill_t
            c.font = hfont(bold=True)
            c.border = thin()
            c.alignment = Alignment(horizontal='center')
            if ci in {4,5,6,8,9} and isinstance(val,(int,float)) and val != '':
                c.number_format = '#,##0.00 "€"'
            if ci == 7 and isinstance(val,(int,float)) and val != '':
                c.number_format = '0.00"%"'

    # ── Lost banner ────────────────────────────────────────────────────────
    if not has_any_post:
        row += 1
        ws.merge_cells(f'A{row}:J{row}')
        c = ws.cell(row, 1, '⚠ VERKEHR / KUNDE VERLOREN – keine POST-Sendungen fakturiert')
        c.font = hfont(bold=True, size=12, color='FFFFFF')
        c.fill = RED_FILL
        c.alignment = Alignment(horizontal='center')
        ws.row_dimensions[row].height = 22

    # ── Section B: PRE samples ─────────────────────────────────────────────
    row += 2
    cust_samples_pre = samples_df[
        (samples_df['Kunden Nr BK'] == knr) & (samples_df['periode'] == 'PRE')
    ]
    cust_samples_post = samples_df[
        (samples_df['Kunden Nr BK'] == knr) & (samples_df['periode'] == 'POST')
    ]

    for section_label, sect_fill, sect_samples in [
        ('B — Repräsentative PRE-Sendungen (nächste am Median-Erlös)', PRE_FILL, cust_samples_pre),
        ('C — Repräsentative POST-Sendungen (nächste am Median-Erlös)', POST_FILL, cust_samples_post),
    ]:
        n_cols = len(SHOW_COLS)
        end_col = get_column_letter(n_cols)
        ws.merge_cells(f'A{row}:{end_col}{row}')
        c = ws.cell(row, 1, section_label)
        c.font = hfont(bold=True, size=11, color='FFFFFF')
        c.fill = LBLUE_FILL
        ws.row_dimensions[row].height = 20

        row += 1
        for ci, col_name in enumerate(SHOW_COLS, 1):
            c = ws.cell(row, ci, col_name)
            c.font = hfont(bold=True)
            c.fill = sect_fill
            c.border = thin()
            c.alignment = Alignment(horizontal='center', wrap_text=True)
            set_width(ws, ci, COL_WIDTHS.get(col_name, 14))
        ws.row_dimensions[row].height = 28

        if len(sect_samples) == 0:
            row += 1
            ws.cell(row, 1, 'Keine fakturierten Sendungen in diesem Zeitraum.').font = hfont(bold=True)
        else:
            for _, srow in sect_samples.iterrows():
                row += 1
                for ci, col_name in enumerate(SHOW_COLS, 1):
                    val = srow.get(col_name, '')
                    if pd.isna(val) if not isinstance(val, str) else False:
                        val = ''
                    # convert timestamps to date strings
                    if hasattr(val, 'date'):
                        val = str(val.date())
                    c = ws.cell(row, ci, val)
                    c.border = thin()
                    c.alignment = Alignment(horizontal='center')
                    if col_name in NUM_COLS and isinstance(val, (int, float)) and val != '':
                        if col_name == 'DB II%':
                            c.number_format = '0.00"%"'
                        else:
                            c.number_format = '#,##0.00'
        row += 2  # gap before next section

print(f"Added {len(top20_meta)} customer sheets")
wb.save(str(XLSX))
print(f"Saved: {XLSX}")
