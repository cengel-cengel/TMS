"""Shared helpers for Excel scope sheets."""
import numpy as np
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BLUE  = "1F4E79"
LGRAY = "D9D9D9"
GRN   = "C6EFCE"; GRN_F = "375623"
RED   = "FFC7CE"; RED_F = "C00000"
YLW   = "FFEB9C"; YLW_F = "7F6000"
DGRAY = "808080"

SCOPE_COLS = [
    ("Empf. Land","land",8),
    ("Empf. PLZ","plz",10),
    ("Ton/Stp/LDM","basis",9),
    ("Erlöse Fracht AX","ef",14),
    ("Erlöse Maut AX","maut",12),
    ("Erlöse Diesel AX","diesel",12),
    ("DLV-Soll","dlv",14),
    ("Δ Fracht","delta",12),
    ("Δ %","fp_pct",9),
    ("M-Klasse","mclass",9),
    ("Hinweis","notes",20),
]

def mc(fp):
    if fp is None or (isinstance(fp,float) and np.isnan(fp)): return None
    return "M1" if abs(fp)<=0.05 else ("M2" if fp<-0.05 else "M_over")

def write_scope_sheet(wb, sheet_name, rows, meta=""):
    """Write a scope sheet. rows = list of dicts with keys from SCOPE_COLS[*][1]."""
    ws = wb.create_sheet(title=sheet_name[:31])
    # Meta line
    ws.merge_cells(f"A1:{get_column_letter(len(SCOPE_COLS))}1")
    ws["A1"].value = meta
    ws["A1"].font = Font(name="Calibri", italic=True, size=9)
    ws.row_dimensions[1].height = 14

    # Header
    for c, (hdr, key, w) in enumerate(SCOPE_COLS, 1):
        cell = ws.cell(row=2, column=c, value=hdr)
        cell.font = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[2].height = 24

    # Data rows
    for r, row in enumerate(rows, 3):
        mc_val = row.get("mclass","")
        for c, (hdr, key, w) in enumerate(SCOPE_COLS, 1):
            v = row.get(key,"")
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name="Calibri", size=9)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if key in ("ef","maut","diesel","dlv","delta"):
                if v != "" and v is not None:
                    try:
                        cell.number_format = '#,##0.00'
                    except: pass
            if key == "fp_pct" and v != "":
                try:
                    cell.number_format = '0.0%'
                except: pass
            if key == "mclass":
                if v == "M1":
                    cell.fill = PatternFill("solid", fgColor=GRN)
                    cell.font = Font(name="Calibri", size=9, color=GRN_F, bold=True)
                elif v == "M2":
                    cell.fill = PatternFill("solid", fgColor=RED)
                    cell.font = Font(name="Calibri", size=9, color=RED_F, bold=True)
                elif v == "M_over":
                    cell.fill = PatternFill("solid", fgColor=YLW)
                    cell.font = Font(name="Calibri", size=9, color=YLW_F, bold=True)
                elif v in ("dlv_luecke","ton_zero","OOS"):
                    cell.fill = PatternFill("solid", fgColor=DGRAY)

    # Totals
    tr = len(rows)+3
    ws.cell(row=tr, column=1, value="SUMME").font = Font(name="Calibri",bold=True,size=10)
    tot_ef = sum(r.get("ef",0) or 0 for r in rows)
    tot_maut = sum(r.get("maut",0) or 0 for r in rows)
    tot_dlv = sum(r.get("dlv",0) or 0 for r in rows)
    tot_delta = sum(r.get("delta",0) or 0 for r in rows)
    for cidx, val in [(4,tot_ef),(5,tot_maut),(7,tot_dlv),(8,tot_delta)]:
        c2 = ws.cell(row=tr, column=cidx, value=round(val,2))
        c2.font = Font(name="Calibri", bold=True, size=10)
        c2.number_format = '#,##0.00'

    # Freeze + filter
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(SCOPE_COLS))}{len(rows)+2}"
    return ws
