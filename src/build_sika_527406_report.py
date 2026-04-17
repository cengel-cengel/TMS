#!/usr/bin/env python3
"""build_sika_527406_report.py — Sika Automotive AG (KNR 527406), gewichtsbasiert"""
import math, re
from pathlib import Path
import numpy as np, pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NK_XLSX  = Path('data/extracted/v1/Noerpel AI/SIka/Nebenbedingungen DINAS/NK Sika.xlsx')
DLV_DIR  = Path('data/extracted/v1/Noerpel AI/SIka/DLV/SIKA Automotive/2026/Upload')
BI_XLSX  = Path('data/bi_report/Tagesbericht.Einzeldaten.alle.VKA.5.xlsx')
BI_CACHE = Path('output/bi_cache_sika_527406.pkl')
OUT_XLSX = Path('output/billing_report/sika_527406_dinas_vergleich.xlsx')
OUT_XLSX.parent.mkdir(exist_ok=True)
KNR, KUNDE, BASIS = 527406, 'Sika Automotive AG', 'EUR/100kg'
MIGRATION_DATE = pd.Timestamp('2025-09-26')

KG_BANDS = [300, 500, 1000, 2000, 5000, 10000, 20000]

def kg_band(n):
    try: n = float(n)
    except: return '?'
    if math.isnan(n) or n <= 0: return '?'
    for b in KG_BANDS:
        if n <= b: return f'bis {b} kg'
    return 'ueber 20000 kg'

def billing_kg(toneff):
    try:
        t = float(toneff)
        if math.isnan(t) or t <= 0: return None
        return math.ceil(t / 100) * 100
    except: return None

# ── DLV Sika Automotive Gewichts-Tarif ──────────────────────────────────────
def load_sika_atm_dlv():
    rates = {}
    for f in sorted(DLV_DIR.glob('V_FRA_7042_O_*_Sika_Automotive.xlsx')):
        df = pd.read_excel(f, header=None)
        bis_weights = []
        for v in df.iloc[8, 1:]:
            if pd.isna(v): break
            try: bis_weights.append(int(float(v)))
            except: break
        for _, row in df.iloc[11:].iterrows():
            zone = str(row.iloc[0]).strip()
            if not zone or zone == 'nan' or pd.isna(row.iloc[0]): break
            prices = {}
            for i, bis in enumerate(bis_weights, 1):
                v = row.iloc[i]
                if not pd.isna(v):
                    try: prices[bis] = float(v)
                    except: pass
            if prices:
                rates[zone] = prices
    return rates

def _zone_from_plz(land, plz):
    plz = str(plz).strip().upper()
    for pfx in ('GB-','NL-','BE-','PT-','FR-','ES-','IT-','AT-','RS-'):
        plz = plz.replace(pfx, '').strip()
    land = str(land).strip().upper()
    if land == 'GB':
        m = re.match(r'([A-Z]+)', plz)
        return f'GB{m.group(1)}' if m else None
    m = re.match(r'(\d{2})', plz) or re.match(r'(\d)', plz)
    if not m: return None
    prefix = m.group(1)
    n = int(prefix)
    if land == 'FR': return f'FR{n}'
    if land == 'ES': return f'ES{n}'
    if land == 'IT': return f'IT{n}'
    if land in ('AT','BE','NL','PT','RS'): return f'{land}{prefix}'
    return None

def lookup_sika_atm(dlv, land, plz, toneff):
    bkg = billing_kg(toneff)
    if bkg is None: return None
    zone = _zone_from_plz(land, plz)
    if zone is None or zone not in dlv: return None
    band_prices = dlv[zone]
    for bis_kg in sorted(band_prices.keys()):
        if bis_kg >= bkg:
            return band_prices[bis_kg]
    return band_prices[max(band_prices.keys())]

print('Lade Sika Automotive DLV...')
dlv = load_sika_atm_dlv()
print(f'DLV: {len(dlv)} Zonen')

# ── BI laden & splitten ─────────────────────────────────────────────────────
if BI_CACHE.exists():
    df_knr = pd.read_pickle(BI_CACHE)
    print(f'BI Cache geladen: {len(df_knr)} Zeilen')
else:
    print('Lade BI Daten (kann etwas dauern)...')
    df_all = pd.read_excel(BI_XLSX, header=0)
    df_knr = df_all[df_all['Kunden Nr BK'] == KNR].copy()
    df_knr.to_pickle(BI_CACHE)
    print(f'KNR {KNR}: {len(df_knr)} Zeilen, Cache gespeichert')

df_knr['Leistungsdatum'] = pd.to_datetime(df_knr['Leistungsdatum'], errors='coerce')

ERLOESE_MAP = {
    'Erlöse Fracht': 'Fracht_raw',
    'Erlöse Diesel': 'Diesel_raw',
    'Erlöse Maut':   'Maut_raw',
    'Erlöse Nebengebühr': 'Neben_raw',
    'Erlöse Lademittel':  'Lademittel_raw',
    'Erlöse Peak':        'Peak_raw',
    'Erlöse EUST Zoll':   'EUST_raw',
    'Erlöse Transportversicherung': 'Versich_raw',
    'Erloese': 'Gesamt_raw',
}
for old, new in ERLOESE_MAP.items():
    if old in df_knr.columns:
        df_knr[new] = pd.to_numeric(df_knr[old], errors='coerce')

pre_mask  = df_knr['Leistungsdatum'] < MIGRATION_DATE
post_mask = df_knr['Leistungsdatum'] >= MIGRATION_DATE
pre  = df_knr[pre_mask].copy()
post = df_knr[post_mask].copy()
print(f'PRE: {len(pre)} Zeilen, POST: {len(post)} Zeilen')

PRE_RENAME = {
    'Fracht_raw':'Dinas Fracht',  'Diesel_raw':'Dinas Diesel',
    'Maut_raw':'Dinas Maut/SSD',  'Neben_raw':'Dinas Nebengebühr',
    'Lademittel_raw':'Dinas Lademittel', 'Peak_raw':'Dinas Peak',
    'EUST_raw':'Dinas EUST/Zoll', 'Versich_raw':'Dinas Versicherung',
    'Gesamt_raw':'Dinas Gesamt',
}
POST_RENAME = {
    'Fracht_raw':'AX Fracht',     'Diesel_raw':'AX Diesel',
    'Maut_raw':'AX Maut',         'Neben_raw':'AX Nebengebühr',
    'Lademittel_raw':'AX Lademittel', 'Peak_raw':'AX Peak',
    'EUST_raw':'AX EUST/Zoll',    'Versich_raw':'AX Versicherung',
    'Gesamt_raw':'AX Gesamt',
}
pre  = pre.rename(columns=PRE_RENAME)
post = post.rename(columns=POST_RENAME)

for df, name in [(pre,'PRE'),(post,'POST')]:
    bad = df['Rechnungsnummer'].apply(
        lambda v: str(v).strip().rstrip('0').rstrip('.') in ('0','','nan'))
    if bad.sum(): print(f'Filtere {bad.sum()} {name} Zeilen RN=0')
    df.drop(index=df[bad].index, inplace=True)

for c in list(POST_RENAME.values()):
    if c in post.columns: post[c] = pd.to_numeric(post[c], errors='coerce')

pre = pre[pre['Dinas Fracht'].fillna(0) > 0].copy()
print(f'PRE nach Filter: {len(pre)} Zeilen, POST: {len(post)} Zeilen')

def add_empty_master_cols(df):
    df = df.copy()
    for c in ('_master_nr','_sub_nrs','_ist_master'): df[c] = None
    df['_n_subs'] = 0
    return df

pre  = add_empty_master_cols(pre)
post = add_empty_master_cols(post)

for df in (pre, post):
    df['_bkg']   = df['Tonnage (eff.)'].apply(billing_kg)
    df['_land']  = df['Empfänger Land'].astype(str).str.strip()
    df['_plz2']  = df['Empfänger PLZ'].astype(str).str.strip().str[:2]
    df['_vplz2'] = df['Versender PLZ'].astype(str).str.strip().str[:2]
    df['_kb']    = df['Tonnage (eff.)'].apply(kg_band)
    df['_cl']    = df['_land'] + '|' + df['_plz2'] + '|' + df['_kb']

pre['_eff']  = pre['Dinas Fracht']  / (pre['_bkg'].replace(0, np.nan) / 100)
post['_eff'] = post['AX Fracht']    / (post['_bkg'].replace(0, np.nan) / 100)

print('Berechne Soll EUR...')
pre['Soll EUR']  = pre.apply(
    lambda r: lookup_sika_atm(dlv, r['_land'], r['Empfänger PLZ'], r['Tonnage (eff.)']), axis=1)
post['Soll EUR'] = post.apply(
    lambda r: lookup_sika_atm(dlv, r['_land'], r['Empfänger PLZ'], r['Tonnage (eff.)']), axis=1)
print(f'Soll EUR: PRE {pre["Soll EUR"].notna().sum()}/{len(pre)}, POST {post["Soll EUR"].notna().sum()}/{len(post)}')

# ── Cluster-Statistiken ─────────────────────────────────────────────────────
pa = pre.groupby('_cl').agg(
    n_pre=('Dinas Gesamt','count'),    avg_d=('Dinas Gesamt','mean'),
    avg_df=('Dinas Fracht','mean'),    avg_dd=('Dinas Diesel','mean'),
    avg_dm=('Dinas Maut/SSD','mean'),  vplz2=('_vplz2','first'),
    avg_eff_d=('_eff','mean'),         avg_dlv=('Soll EUR','mean')).reset_index()
oa = post.groupby('_cl').agg(
    n_post=('AX Gesamt','count'),  avg_ax=('AX Gesamt','mean'),
    avg_af=('AX Fracht','mean'),   avg_ad=('AX Diesel','mean'),
    avg_am=('AX Maut','mean'),     avg_eff_ax=('_eff','mean')).reset_index()
stats = pa.merge(oa, on='_cl', how='inner')
stats = stats[stats['avg_ax'] < stats['avg_d']].copy()
stats['delta']        = stats['avg_ax'] - stats['avg_d']
stats['pct']          = stats['delta'] / stats['avg_d'].replace(0, np.nan) * 100
stats['loss']         = stats['delta'] * stats['n_post']
stats['delta_eff_pct'] = ((stats['avg_eff_ax'] - stats['avg_eff_d'])
                           / stats['avg_eff_d'].replace(0, np.nan) * 100)
stats = stats[stats['delta_eff_pct'].abs() > 5].copy()
stats = stats.sort_values('loss').reset_index(drop=True)
print(f'Cluster (|ΔEff|>5%): {len(stats)}, Sigma Verlust: {stats["loss"].sum():,.0f} EUR')

# ── NK-Mapping ──────────────────────────────────────────────────────────────
def get_nk_alt(r):
    return (r.get('Dinas Fracht') or 0,    r.get('Dinas Diesel') or 0,
            r.get('Dinas Maut/SSD') or 0,  r.get('Dinas Lademittel') or 0,
            r.get('Dinas Peak') or 0,       r.get('Dinas Nebengebühr') or 0,
            r.get('Dinas EUST/Zoll') or 0,  r.get('Dinas Versicherung') or 0)

def get_nk_neu(r):
    return (r.get('AX Fracht') or 0,       r.get('AX Diesel') or 0,
            r.get('AX Maut') or 0,          r.get('AX Lademittel') or 0,
            r.get('AX Peak') or 0,          r.get('AX Nebengebühr') or 0,
            r.get('AX EUST/Zoll') or 0,     r.get('AX Versicherung') or 0)

def abw_grund_row(nk, soll, erloese):
    if pd.isna(soll) or soll == 0: return 'Soll n/a'
    pct = (erloese - soll) / soll * 100
    if abs(pct) < 5: return 'OK'
    return 'Unterfakturierung' if pct < 0 else 'Überfakturierung'

# ── Styles ──────────────────────────────────────────────────────────────────
def fill(h): return PatternFill('solid', fgColor=h)
def fnt(bold=False, color='000000', size=9):
    return Font(bold=bold, color=color, size=size)
THIN = Side(border_style='thin', color='BBBBBB')
BRD  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
EUR_FMT = '#,##0.00'
HDR_COLS = ['System','Auftrags-Nr','Master-Nr','Sub-Nr(n)','Anzahl Subs','Ist Master',
            'Rech.-Nr','Sendungsdatum','Kunde','Land','Empf.PLZ','Vers.PLZ',
            'Gew.band','Zone','Basis','Basis Menge','Basispreis',
            'Eff. Preis','Tonnage kg','Stellplätze','Lademeter','Volumen','Soll EUR',
            'Fracht EUR','Diesel EUR','Maut EUR','Lademittel','Peak EUR',
            'Neben EUR','EUST Zoll','Versich.',
            'Erlöse','Abw. Grund']
N = len(HDR_COLS)
EUR_COLS  = {17,18,23,24,25,26,27,28,29,30,31,32}
NUM_RIGHT = {5, 16, 19, 20, 21, 22}
STR_COLS  = {2, 3, 7}; DATE_COL = 8
FILL_HDR  = fill('1F497D'); FILL_CLU = fill('2E75B6')
FILL_ALT  = fill('BDD7EE'); FILL_NEU = fill('FCE4D6'); FILL_CTRL = fill('E2EFDA')

def wc(ws, r, c, v=None, f=None, fn=None, al='left', fmt=None, brd=None):
    cell = ws.cell(row=r, column=c)
    if v is not None: cell.value = v
    if f:   cell.fill = f
    if fn:  cell.font = fn
    if fmt: cell.number_format = fmt
    if brd: cell.border = brd
    cell.alignment = Alignment(horizontal=al, vertical='center')

def write_row(ws, row, values, row_fill, is_ctrl=False):
    rf = FILL_CTRL if is_ctrl else row_fill
    for ci, v in enumerate(values, 1):
        if v is not None and isinstance(v, float) and math.isnan(v): v = None
        if ci in STR_COLS and v is not None:
            try: v = str(int(float(str(v))))
            except: v = str(v)
        fmt = ('DD.MM.YYYY' if ci == DATE_COL else EUR_FMT if ci in EUR_COLS else None)
        al  = 'right' if ci in EUR_COLS or ci in NUM_RIGHT else 'left'
        wc(ws, row, ci, v, rf, fnt(size=9), al, fmt, BRD)

# ── Sheet 1: Hauptauswertung ─────────────────────────────────────────────────
def build_main_sheet(ws):
    ws.title = 'Sika Automotive AG'
    ws.freeze_panes = 'A3'
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=N)
    wc(ws,1,1, f'Sika Automotive AG ({KNR}) — Dinas PRE vs AX POST  |  '
               f'Cluster: {len(stats)}  |  Sigma Verlust: {stats["loss"].sum():,.0f} EUR',
       FILL_HDR, fnt(bold=True, color='FFFFFF', size=12), 'center')
    for ci, h in enumerate(HDR_COLS, 1):
        wc(ws,2,ci,h,FILL_HDR,fnt(bold=True,color='FFFFFF',size=9),'center',brd=BRD)
    row = 2
    for _, cl in stats.iterrows():
        ckey = cl['_cl']; parts = ckey.split('|')
        land, plz_p, kb = (parts+['','',''])[:3]
        vplz = str(cl.get('vplz2',''))
        pre_all  = pre[pre['_cl']==ckey]
        post_all = post[post['_cl']==ckey]
        post_under = post_all[post_all['AX Fracht'].fillna(float('inf')) < cl.avg_df]
        post_under = post_under[[
            abw_grund_row(get_nk_neu(r), r.get('Soll EUR'), r.get('AX Gesamt') or 0) != 'Überfakturierung'
            for _, r in post_under.iterrows()
        ]]
        if len(post_under) == 0: continue
        pre_s = pre_all.head(5); post_s = post_under.head(5)
        ctrl_pi = ctrl_oi = None; best = float('inf')
        for pi, pr in pre_all.iterrows():
            for oi, po in post_under.iterrows():
                d = abs(float(pr.get('_bkg') or 0) - float(po.get('_bkg') or 0))
                if d < best: best=d; ctrl_pi=pi; ctrl_oi=oi
        row += 1
        dlv_str = f'{cl.avg_dlv:,.2f}' if not pd.isna(cl.avg_dlv) else 'n/a'
        lbl = (f'▶ {KNR}|{vplz}|{land}|{plz_p}|{kb}|{BASIS}     '
               f'n_PRE={int(cl.n_pre)}  n_POST={int(cl.n_post)}  '
               f'Ø Dinas={cl.avg_d:,.2f} EUR  Ø AX={cl.avg_ax:,.2f} EUR  '
               f'Δ={cl.delta:,.2f} EUR ({cl.pct:+.1f}%)  '
               f'est.Verlust={cl.loss:,.0f} EUR  |  '
               f'Eff. Dinas={cl.avg_eff_d:,.2f} | Eff. AX={cl.avg_eff_ax:,.2f} | '
               f'DLV={dlv_str} | ΔEff={cl.delta_eff_pct:+.1f}%')
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=N)
        wc(ws, row, 1, lbl, FILL_CLU, fnt(bold=True, color='FFFFFF', size=10))
        for _, r in pre_s.iterrows():
            row += 1
            nk = get_nk_alt(r); erloese = r.get('Dinas Gesamt') or 0
            soll = r.get('Soll EUR'); bkg = r.get('_bkg')
            eff = r.get('_eff')
            bp = (soll / (bkg / 100)) if soll and bkg else None
            zone = _zone_from_plz(r.get('_land',''), r.get('Empfänger PLZ',''))
            vals = ['alt', r.get('Auftragsnummer'),
                    r.get('_master_nr'), r.get('_sub_nrs'), int(r.get('_n_subs') or 0), r.get('_ist_master') or '',
                    r.get('Rechnungsnummer'), r.get('Leistungsdatum'), KUNDE,
                    r.get('Empfänger Land'), r.get('Empfänger PLZ'), r.get('Versender PLZ'),
                    kb, zone, BASIS, bkg, bp, eff,
                    r.get('Tonnage (eff.)'), r.get('Stellplätze'), r.get('Lademeter'), r.get('Volumen'), soll,
                    *nk, erloese, abw_grund_row(nk, soll, erloese)]
            write_row(ws, row, vals, FILL_ALT, r.name==ctrl_pi)
        for _, r in post_s.iterrows():
            row += 1
            nk = get_nk_neu(r); erloese = r.get('AX Gesamt') or 0
            soll = r.get('Soll EUR'); bkg = r.get('_bkg')
            eff = r.get('_eff')
            bp = (soll / (bkg / 100)) if soll and bkg else None
            zone = _zone_from_plz(r.get('_land',''), r.get('Empfänger PLZ',''))
            vals = ['neu', r.get('Auftragsnummer'),
                    r.get('_master_nr'), r.get('_sub_nrs'), int(r.get('_n_subs') or 0), r.get('_ist_master') or '',
                    r.get('Rechnungsnummer'), r.get('Leistungsdatum'), KUNDE,
                    r.get('Empfänger Land'), r.get('Empfänger PLZ'), r.get('Versender PLZ'),
                    kb, zone, BASIS, bkg, bp, eff,
                    r.get('Tonnage (eff.)'), r.get('Stellplätze'), r.get('Lademeter'), r.get('Volumen'), soll,
                    *nk, erloese, abw_grund_row(nk, soll, erloese)]
            write_row(ws, row, vals, FILL_NEU, r.name==ctrl_oi)
        row += 1
    widths = [8,15,16,30,8,9, 13,12,18,5,8,8, 14,8,10, 9,10,10, 9,9,9,9, 10, 10,9,9,9,9,9,9,9, 11,22]
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[1].height = 20; ws.row_dimensions[2].height = 18
    return row

def build_nk_sheet(ws):
    ws.title = 'NK_Konditionen'
    nk_wb = load_workbook(NK_XLSX, data_only=True); nk_ws = nk_wb.active
    for r_idx, r in enumerate(nk_ws.iter_rows(values_only=True), 1):
        for c_idx, val in enumerate(r, 1):
            ws.cell(row=r_idx, column=c_idx).value = val
    for ci in range(1, (nk_ws.max_column or 19)+1):
        ws.column_dimensions[get_column_letter(ci)].width = 22

wb = Workbook()
last_row = build_main_sheet(wb.active)
build_nk_sheet(wb.create_sheet('NK_Konditionen'))
wb.save(OUT_XLSX)
print(f'\nGespeichert: {OUT_XLSX}')
print(f'  Sheet "Sika Automotive AG": {last_row} Zeilen')
print(f'  Sheet "NK_Konditionen": NK Sika.xlsx kopiert')
