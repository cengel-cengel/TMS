"""Word Skript — Konversion FINAL_AUDIT_REPORT_v2_0.md -> Audit_Bericht_v2_0.docx."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = Path("docs/FINAL_AUDIT_REPORT_v2_0.md")
OUT = Path("output/erka_lieferung/Audit_Bericht_v2_0.docx")

doc = Document()

# Page margins
for sec in doc.sections:
    sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)

# Default style
style = doc.styles['Normal']
style.font.name = 'Calibri'; style.font.size = Pt(11)

# Header/Footer
sec = doc.sections[0]
hdr = sec.header; hdr.is_linked_to_previous = False
hp = hdr.paragraphs[0]
hp.text = "TMS-Migrations-Audit Dinas → AX  |  ERKA-Management  |  2026-04-29"
for run in hp.runs:
    run.font.size = Pt(9); run.font.color.rgb = RGBColor(0x40,0x40,0x40)
ftr = sec.footer
fp2 = ftr.paragraphs[0]
fp2.text = "Abschlussbericht v2.0  |  Vertraulich  |  Seite "
for run in fp2.runs:
    run.font.size = Pt(9)

def add_page_break(doc):
    doc.add_page_break()

def apply_heading(para, level):
    BLUE = RGBColor(0x1F,0x4E,0x79)
    sizes = {1:16, 2:14, 3:12, 4:11}
    para.style = doc.styles[f'Heading {min(level,4)}']
    for run in para.runs:
        run.font.name = 'Calibri'
        run.font.size = Pt(sizes.get(level,11))
        run.font.color.rgb = BLUE
        run.bold = True

def parse_table(lines):
    """Parse markdown table lines into list of lists."""
    rows = []
    for line in lines:
        if re.match(r'\s*\|[-:| ]+\|\s*$', line): continue
        cells = [c.strip().strip('*') for c in line.strip().strip('|').split('|')]
        rows.append(cells)
    return rows

lines = SRC.read_text(encoding='utf-8').splitlines()
i = 0
page_break_headings = {'§2','§3','§4','§5','§6','§7','§8','§9','§10'}
in_table = False
table_lines = []

# Title page
t = doc.add_heading('Migrationsaudit Dinas → AX', 0)
t.runs[0].font.size = Pt(20); t.runs[0].font.color.rgb = RGBColor(0x1F,0x4E,0x79)
doc.add_paragraph('Abschlussbericht v2.0').runs[0].font.size = Pt(14)
doc.add_paragraph('ERKA-Management / Noerpel-Gruppe').runs[0].italic = True
doc.add_paragraph('Stand: 2026-04-29  |  Methodik: v1.9.7').runs[0].font.size = Pt(11)
doc.add_page_break()

while i < len(lines):
    line = lines[i]

    # Detect table block
    if line.startswith('|') and not in_table:
        in_table = True
        table_lines = [line]
        i += 1
        continue
    if in_table:
        if line.startswith('|'):
            table_lines.append(line)
            i += 1
            continue
        else:
            # Flush table
            rows = parse_table(table_lines)
            if rows:
                tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
                tbl.style = 'Table Grid'
                for ri, row in enumerate(rows):
                    for ci, cell_text in enumerate(row):
                        if ci < len(tbl.rows[ri].cells):
                            c = tbl.rows[ri].cells[ci]
                            c.text = cell_text
                            c.paragraphs[0].runs[0].font.size = Pt(9) if ri>0 else Pt(10)
                            if ri == 0:
                                c.paragraphs[0].runs[0].bold = True
            table_lines = []; in_table = False
            continue

    # Headings
    m = re.match(r'^(#{1,4})\s+(.*)', line)
    if m:
        level = len(m.group(1))
        text = m.group(2).strip().lstrip('*').rstrip('*')
        # Page break before major sections
        if level <= 2 and any(s in text for s in page_break_headings):
            doc.add_page_break()
        para = doc.add_heading(text, level=min(level,4))
        apply_heading(para, level)
        i += 1
        continue

    # Horizontal rule
    if line.strip().startswith('---'):
        i += 1
        continue

    # Code block
    if line.strip().startswith('```'):
        i += 1
        code_lines = []
        while i < len(lines) and not lines[i].strip().startswith('```'):
            code_lines.append(lines[i])
            i += 1
        if code_lines:
            p = doc.add_paragraph('\n'.join(code_lines))
            p.style = 'Normal'
            p.runs[0].font.name = 'Courier New'
            p.runs[0].font.size = Pt(9)
        i += 1
        continue

    # Blockquote
    if line.strip().startswith('>'):
        text = line.strip().lstrip('>').strip()
        if text:
            p = doc.add_paragraph(text)
            p.paragraph_format.left_indent = Cm(1)
            p.runs[0].italic = True
        i += 1
        continue

    # Bullet
    if re.match(r'^[\*\-]\s+', line.strip()):
        text = re.sub(r'^[\*\-]\s+', '', line.strip())
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        p = doc.add_paragraph(text, style='List Bullet')
        p.runs[0].font.size = Pt(10)
        i += 1
        continue

    # Empty line
    if not line.strip():
        i += 1
        continue

    # Normal paragraph — strip markdown bold
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', line)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    if text.strip():
        p = doc.add_paragraph(text.strip())
        p.runs[0].font.size = Pt(11) if not text.startswith('    ') else Pt(9)
    i += 1

doc.save(OUT)
print(f"Saved: {OUT}  ({len(doc.paragraphs)} paragraphs)")
