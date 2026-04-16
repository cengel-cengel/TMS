#!/usr/bin/env python3
"""build_geze_report.py — GEZE GmbH (KNR 406035), Tonnage-basiert EUR/100kg"""
import glob as _glob, math, re
from pathlib import Path
import numpy as np, pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import sys; sys.path.insert(0, 'src')
from dinas_pdf_parser import parse_one, flatten

SRC_XLSX  = Path('output/geze_dinas_vergleich.xlsx')
NK_XLSX   = Path('data/extracted/v1/Noerpel AI/GEZE/Nebenbedingungen DINAS/NK Geze.xlsx')
DLV_PATH  = Path('data/extracted/v1/Noerpel AI/GEZE/DLV/2025/20250305_Geze_Export_incl. MP_ PT_GB_IT_FR_AT_ES_CH_inkl. Maut und Zusatzkosten IT-00_ERG#U00c4NZT UM DUBLIN.xlsx')
DINAS_DIR = Path('data/extracted/v1/Noerpel AI/GEZE/Rechnungen/Rechnungen DINAS')
CACHE     = Path('output/dinas_cache_406035.pkl')
OUT_XLSX  = Path('output/billing_report/geze_dinas_vergleich.xlsx')
OUT_XLSX.parent.mkdir(exist_ok=True)
KNR, KUNDE, BASIS = 406035, 'GEZE GmbH', 'EUR/100kg'

BANDS = [100, 200, 300, 500, 750, 1000, 2000, 3000]
def gew_band(kg):
    try: kg = float(kg)
    except: return '?'
    if math.isnan(kg) or kg <= 0: return '?'
    for b in BANDS:
        if kg <= b: return f'bis {b}kg'
    return 'ueber 3000kg'

# ── DLV GEZE (Exporttarife, weight-based per 100kg) ───────────────────────
COUNTRY_MAP = {
    'Portugal': 'PT', 'Großbritanien': 'GB', 'Irland': 'IE',
    'Italien': 'IT', 'Frankreich': 'FR', 'Frankreich Zone 7': 'FR',
    'Österreich': 'AT', 'Spanien': 'ES', 'Schweiz': 'CH',
}

def _extract_prefixes(plz_str, land):
    prefixes = set()
    for part in re.split(r'[,;]', str(plz_str)):
        part = part.strip()
        if not part or part == 'nan': continue
        m = re.match(r'^(\d{2})\s*[-–]\s*(\d{2})$', part.replace(' ', ''))
        if m:
            for i in range(int(m.group(1)), int(m.group(2)) + 1):
                prefixes.add(f'{i:02d}')
            continue
        nums = re.findall(r'\b(\d{2})\b', part)
        prefixes.update(nums)
        if land == 'GB':
            m2 = re.match(r'^([A-Z]{1,2})', part.strip().upper())
            if m2: prefixes.add(m2.group(1))
    return prefixes

def _parse_kg(s):
    try: return int(re.sub(r'[^\d]', '', str(s).split('kg')[0]))
    except: return None

def load_geze_dlv():
    df = pd.read_excel(DLV_PATH, sheet_name='Exporttarife', header=None)
    result = {}
    cur = None
    band_cols = {}
    for i, row in df.iterrows():
        v0 = str(row.iloc[0]).strip()
        if v0 in ('nan', '', 'NaN'): continue
        if v0 in COUNTRY_MAP:
            cur = COUNTRY_MAP[v0]
            if cur not in result: result[cur] = []
            band_cols = {}
        elif v0.startswith('ab Werk'):
            for ci, cell in enumerate(row):
                kg = _parse_kg(cell) if 'kg' in str(cell) else None
                if kg: band_cols[ci] = kg
        elif cur and re.match(r'^Zone\s*\d', v0):
            plz_str = str(row.iloc[1])
            try: minimum = float(row.iloc[2])
            except: minimum = 0.0
            prefixes = _extract_prefixes(plz_str, cur)
            zone_bands = []
            for ci, kg in sorted(band_cols.items()):
                try: zone_bands.append((kg, float(row.iloc[ci])))
                except: pass
            if prefixes and zone_bands:
                result[cur].append((v0, prefixes, minimum, zone_bands))
    return result

def lookup_geze(dlv, land, plz, tonnage_kg):
    zones = dlv.get(land, [])
    if not zones or pd.isna(tonnage_kg): return None
    plz = str(plz).strip().upper()
    billing_kg = max(100, math.ceil(float(tonnage_kg) / 100) * 100)
    for zone_name, prefixes, minimum, bands in zones:
        if plz[:2] not in prefixes and plz[:1] not in prefixes: continue
        rate = next((r for bis_kg, r in sorted(bands) if billing_kg <= bis_kg), bands[-1][1])
        return max(minimum, rate * billing_kg / 100)
    return None

print('Lade GEZE DLV...')
dlv = load_geze_dlv()
print(f'DLV: {sum(len(v) for v in dlv.values())} Zonen in {list(dlv.keys())}')

# ── Daten laden ────────────────────────────────────────────────────────────
pre  = pd.read_excel(SRC_XLSX, sheet_name='PRE Dinas Detail',  header=2)
post = pd.read_excel(SRC_XLSX, sheet_name='POST AX Detail',    header=2)

for df, name in [(pre, 'PRE'), (post, 'POST')]:
    bad = df['Rechnungsnummer'].apply(
        lambda v: str(v).strip().rstrip('0').rstrip('.') in ('0', '', 'nan'))
    if bad.sum():
        print(f'Filtere {bad.sum()} {name} Zeilen RN=0')
        df.drop(index=df[bad].index, inplace=True)

POST_NK = ['AX Fracht','AX Diesel','AX Maut','AX Nebengebühr',
           'AX Lademittel','AX Peak','AX EUST/Zoll','AX Versicherung','AX Gesamt']
for c in POST_NK: post[c] = pd.to_numeric(post.get(c), errors='coerce')

# ── DINAS-NK: per-Sendung aus Cache ────────────────────────────────────────
if CACHE.exists():
    dinas_cache = pd.read_pickle(CACHE)
    print(f'DINAS cache: {len(dinas_cache)} Positionen')
else:
    print('Parse DINAS PDFs...')
    pdfs = _glob.glob(str(DINAS_DIR / '*.pdf'))
    rows = [flatten(pos) for p in pdfs for pos in parse_one(p)]
    dinas_cache = pd.DataFrame(rows)
    dinas_cache.to_pickle(CACHE)
    print(f'Cache gespeichert: {len(dinas_cache)} Positionen')

def _norm(v):
    s = re.sub(r'\D', '', str(v)).lstrip('0')
    return s if s else str(v).strip()

dinas_cache['_snr'] = dinas_cache['sendungs_nr'].astype(str).apply(_norm)
pre['_snr'] = pre['Auftragsnummer'].astype(str).apply(_norm)
PRE_NK = ['Dinas Fracht','Dinas Diesel','Dinas Maut/SSD','Dinas Ausfuhr',
          'Dinas Verzollung','Dinas Zoll Duty','Dinas Zollbetrag',
          'Dinas Sulphur','Dinas Nebenkostenpausch.','Dinas Redebit',
          'Dinas Sonstige','Dinas Gesamt']
pre.drop(columns=[c for c in PRE_NK if c in pre.columns], inplace=True)
DINAS_RENAME = {
    'fracht':'Dinas Fracht', 'diesel':'Dinas Diesel', 'maut_ssd':'Dinas Maut/SSD',
    'ausfuhr':'Dinas Ausfuhr', 'verzollung':'Dinas Verzollung',
    'zoll_duty':'Dinas Zoll Duty', 'zoll_betrag':'Dinas Zollbetrag',
    'sulphur':'Dinas Sulphur', 'neben_pausch':'Dinas Nebenkostenpausch.',
    'redebit':'Dinas Redebit', 'sonstige':'Dinas Sonstige', 'gesamtbetrag':'Dinas Gesamt',
}
cache_nk = dinas_cache.rename(columns=DINAS_RENAME)[['_snr'] + list(DINAS_RENAME.values())].copy()
pre = pre.merge(cache_nk, on='_snr', how='left')
pre.drop(columns=['_snr'], inplace=True)
for c in PRE_NK: pre[c] = pd.to_numeric(pre.get(c), errors='coerce')
print(f'DINAS re-joined: {pre["Dinas Fracht"].notna().sum()}/{len(pre)} PRE')
pre = pre[pre['Dinas Fracht'].fillna(0) > 0].copy()

for df in (pre, post):
    df['Tonnage (eff.)'] = pd.to_numeric(df['Tonnage (eff.)'], errors='coerce')
    df['_land']  = df['Empfänger Land'].astype(str).str.strip()
    df['_plz2']  = df['Empfänger PLZ'].astype(str).str.strip().str[:2]
    df['_vplz2'] = df['Versender PLZ'].astype(str).str.strip().str[:2]
    df['_gwb']   = df['Tonnage (eff.)'].apply(gew_band)
    df['_cl']    = df['_land'] + '|' + df['_plz2'] + '|' + df['_gwb']

pre['_eff']  = pre['Dinas Fracht']  / pre['Tonnage (eff.)'] * 100
post['_eff'] = post['AX Fracht']    / post['Tonnage (eff.)'] * 100

# ── Soll EUR via DLV ───────────────────────────────────────────────────────
print('Berechne Soll EUR...')
pre['Soll EUR']  = pre.apply( lambda r: lookup_geze(dlv, r['_land'], r['Empfänger PLZ'], r['Tonnage (eff.)']), axis=1)
post['Soll EUR'] = post.apply(lambda r: lookup_geze(dlv, r['_land'], r['Empfänger PLZ'], r['Tonnage (eff.)']), axis=1)
print(f'Soll EUR: PRE {pre["Soll EUR"].notna().sum()}/{len(pre)}, POST {post["Soll EUR"].notna().sum()}/{len(post)}')

# ── Cluster-Statistiken ────────────────────────────────────────────────────
pa = pre.groupby('_cl').agg(
    n_pre=('Dinas Gesamt','count'), avg_d=('Dinas Gesamt','mean'),
    avg_df=('Dinas Fracht','mean'), avg_dd=('Dinas Diesel','mean'),
    avg_dm=('Dinas Maut/SSD','mean'), vplz2=('_vplz2','first'),
    avg_eff_d=('_eff','mean'), avg_dlv=('Soll EUR','mean')).reset_index()
oa = post.groupby('_cl').agg(
    n_post=('AX Gesamt','count'), avg_ax=('AX Gesamt','mean'),
    avg_af=('AX Fracht','mean'), avg_ad=('AX Diesel','mean'),
    avg_am=('AX Maut','mean'), avg_eff_ax=('_eff','mean')).reset_index()
stats = pa.merge(oa, on='_cl', how='inner')
stats = stats[stats['avg_ax'] < stats['avg_d']].copy()
stats['delta'] = stats['avg_ax'] - stats['avg_d']
stats['pct']   = stats['delta'] / stats['avg_d'].replace(0, np.nan) * 100
stats['loss']  = stats['delta'] * stats['n_post']
stats['delta_eff_pct'] = ((stats['avg_eff_ax'] - stats['avg_eff_d'])
                          / stats['avg_eff_d'].replace(0, np.nan) * 100)
stats = stats[stats['delta_eff_pct'].abs() > 5].copy()
stats = stats.sort_values('loss').reset_index(drop=True)
print(f'Cluster (|ΔEff|>5%): {len(stats)}, Sigma Verlust: {stats["loss"].sum():,.0f} EUR')

# ── NK-Mapping ─────────────────────────────────────────────────────────────
def get_nk_alt(r):
    fracht = r.get('Dinas Fracht') or 0; diesel = r.get('Dinas Diesel') or 0
    maut   = r.get('Dinas Maut/SSD') or 0; neben  = r.get('Dinas Ausfuhr') or 0
    eust   = ((r.get('Dinas Verzollung') or 0) + (r.get('Dinas Zoll Duty') or 0)
               + (r.get('Dinas Zollbetrag') or 0))
    versich = ((r.get('Dinas Sulphur') or 0) + (r.get('Dinas Nebenkostenpausch.') or 0)
               + (r.get('Dinas Redebit') or 0) + (r.get('Dinas Sonstige') or 0))
    return fracht, diesel, maut, 0, 0, neben, eust, versich

def get_nk_neu(r):
    return (r.get('AX Fracht') or 0, r.get('AX Diesel') or 0,
            r.get('AX Maut') or 0, r.get('AX Lademittel') or 0,
            r.get('AX Peak') or 0, r.get('AX Nebengebühr') or 0,
            r.get('AX EUST/Zoll') or 0, r.get('AX Versicherung') or 0)

def abw_grund_row(nk, soll, erloese):
    if pd.isna(soll) or soll == 0: return 'Soll n/a'
    pct = (erloese - soll) / soll * 100
    if abs(pct) < 5: return 'OK'
    return 'Unterfakturierung' if pct < 0 else 'Überfakturierung'

# ── Styles ─────────────────────────────────────────────────────────────────
def fill(h): return PatternFill('solid', fgColor=h)
def fnt(bold=False, color='000000', size=9):
    return Font(bold=bold, color=color, size=size)
THIN = Side(border_style='thin', color='BBBBBB')
BRD  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
EUR_FMT = '#,##0.00'
HDR_COLS = ['System','Auftrags-Nr','Rech.-Nr','Sendungsdatum','Kunde','Land','Empf.PLZ','Vers.PLZ',
            'Gew.band','Zone','Basis','Basis Menge','Basispreis',
            'Eff. Preis','Tonnage kg','Stellplätze','Lademeter','Volumen','Soll EUR',
            'Fracht EUR','Diesel EUR','Maut EUR','Lademittel','Peak EUR',
            'Neben EUR','EUST Zoll','Versich.',
            'Erlöse','Abw. Grund']
N = len(HDR_COLS)
EUR_COLS  = {13,14,19,20,21,22,23,24,25,26,27,28}
NUM_RIGHT = {12, 15, 16, 17, 18}
STR_COLS  = {2, 3}; DATE_COL = 4
FILL_HDR = fill('1F497D'); FILL_CLU = fill('2E75B6')
FILL_ALT = fill('BDD7EE'); FILL_NEU = fill('FCE4D6'); FILL_CTRL = fill('E2EFDA')

def wc(ws, r, c, v=None, f=None, fn=None, al='left', fmt=None, brd=None):
    cell = ws.cell(row=r, column=c)
    if v is not None: cell.value = v
    if f:   cell.fill = f
    if fn:  cell.font = fn
    if fmt: cell.number_format = fmt
    if brd: cell.border = brd
    cell.alignment = Alignment(horizontal=al, vertical='center')

def write_row(ws, row, values, row_fill, is_ctrl=False):
    rf = FILL_CTRL if is_ctrl else row_fill
    for ci, v in enumerate(values, 1):
        if v is not None and isinstance(v, float) and math.isnan(v): v = None
        if ci in STR_COLS and v is not None:
            try: v = str(int(float(str(v))))
            except: v = str(v)
        fmt = ('DD.MM.YYYY' if ci == DATE_COL else EUR_FMT if ci in EUR_COLS else None)
        al  = 'right' if ci in EUR_COLS or ci in NUM_RIGHT else 'left'
        wc(ws, row, ci, v, rf, fnt(size=9), al, fmt, BRD)

# ── Sheet 1: GEZE GmbH ────────────────────────────────────────────────────
def build_main_sheet(ws):
    ws.title = 'GEZE GmbH'
    ws.freeze_panes = 'A3'
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=N)
    wc(ws,1,1, f'GEZE GmbH ({KNR}) — Dinas PRE vs AX POST  |  '
               f'Cluster: {len(stats)}  |  Sigma Verlust: {stats["loss"].sum():,.0f} EUR',
       FILL_HDR, fnt(bold=True, color='FFFFFF', size=12), 'center')
    for ci, h in enumerate(HDR_COLS, 1):
        wc(ws,2,ci,h,FILL_HDR,fnt(bold=True,color='FFFFFF',size=9),'center',brd=BRD)
    row = 2
    for _, cl in stats.iterrows():
        ckey = cl['_cl']; parts = ckey.split('|')
        land, plz_p, gwband = (parts+['','',''])[:3]
        vplz = str(cl.get('vplz2',''))
        pre_all  = pre[pre['_cl']==ckey]
        post_all = post[post['_cl']==ckey]
        post_under = post_all[post_all['AX Fracht'].fillna(float('inf')) < cl.avg_df]
        post_under = post_under[[
            abw_grund_row(get_nk_neu(r), r.get('Soll EUR'), r.get('AX Gesamt') or 0) != 'Überfakturierung'
            for _, r in post_under.iterrows()
        ]]
        if len(post_under) == 0: continue
        pre_s = pre_all.head(5); post_s = post_under.head(5)
        ctrl_pi = ctrl_oi = None; best = float('inf')
        for pi, pr in pre_all.iterrows():
            pr_kg = pr.get('Tonnage (eff.)')
            if pd.isna(pr_kg): continue
            for oi, po in post_under.iterrows():
                po_kg = po.get('Tonnage (eff.)')
                if pd.isna(po_kg): continue
                d = abs(float(pr_kg) - float(po_kg))
                if d < best: best=d; ctrl_pi=pi; ctrl_oi=oi
        row += 1
        dlv_str = f'{cl.avg_dlv:,.2f}' if not pd.isna(cl.avg_dlv) else 'n/a'
        lbl = (f'▶ {KNR}|{vplz}|{land}|{plz_p}|{gwband}|{BASIS}     '
               f'n_PRE={int(cl.n_pre)}  n_POST={int(cl.n_post)}  '
               f'Ø Dinas={cl.avg_d:,.2f} EUR  Ø AX={cl.avg_ax:,.2f} EUR  '
               f'Δ={cl.delta:,.2f} EUR ({cl.pct:+.1f}%)  '
               f'est.Verlust={cl.loss:,.0f} EUR  |  '
               f'Eff. Dinas={cl.avg_eff_d:,.2f} | Eff. AX={cl.avg_eff_ax:,.2f} | '
               f'DLV={dlv_str} | ΔEff={cl.delta_eff_pct:+.1f}%')
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=N)
        wc(ws, row, 1, lbl, FILL_CLU, fnt(bold=True, color='FFFFFF', size=10))
        for _, r in pre_s.iterrows():
            row += 1
            nk = get_nk_alt(r); erloese = r.get('Dinas Gesamt') or 0
            soll = r.get('Soll EUR'); kg = r.get('Tonnage (eff.)'); eff = r.get('_eff')
            billing_kg = math.ceil(float(kg)/100)*100 if soll and pd.notna(kg) else None
            bp = soll / billing_kg * 100 if soll and billing_kg else None
            vals = ['alt', r.get('Auftragsnummer'), r.get('Rechnungsnummer'),
                    r.get('Leistungsdatum'), KUNDE,
                    r.get('Empfänger Land'), r.get('Empfänger PLZ'), r.get('Versender PLZ'),
                    gwband, 'n/a', BASIS, billing_kg, bp, eff,
                    kg, r.get('Stellplätze'), r.get('Lademeter'), r.get('Volumen'), soll,
                    *nk, erloese, abw_grund_row(nk, soll, erloese)]
            write_row(ws, row, vals, FILL_ALT, r.name==ctrl_pi)
        for _, r in post_s.iterrows():
            row += 1
            nk = get_nk_neu(r); erloese = r.get('AX Gesamt') or 0
            soll = r.get('Soll EUR'); kg = r.get('Tonnage (eff.)'); eff = r.get('_eff')
            billing_kg = math.ceil(float(kg)/100)*100 if soll and pd.notna(kg) else None
            bp = soll / billing_kg * 100 if soll and billing_kg else None
            vals = ['neu', r.get('Auftragsnummer'), r.get('Rechnungsnummer'),
                    r.get('Leistungsdatum'), KUNDE,
                    r.get('Empfänger Land'), r.get('Empfänger PLZ'), r.get('Versender PLZ'),
                    gwband, 'n/a', BASIS, billing_kg, bp, eff,
                    kg, r.get('Stellplätze'), r.get('Lademeter'), r.get('Volumen'), soll,
                    *nk, erloese, abw_grund_row(nk, soll, erloese)]
            write_row(ws, row, vals, FILL_NEU, r.name==ctrl_oi)
        row += 1
    widths = [8,15,13,12,18,5,8,8, 11,8,10, 9,10,10,9,9,9,9,10, 10,9,9,9,9,9,9,9, 11,22]
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[1].height = 20; ws.row_dimensions[2].height = 18
    return row

def build_nk_sheet(ws):
    ws.title = 'NK_Konditionen'
    nk_wb = load_workbook(NK_XLSX, data_only=True); nk_ws = nk_wb.active
    for r_idx, r in enumerate(nk_ws.iter_rows(values_only=True), 1):
        for c_idx, val in enumerate(r, 1):
            ws.cell(row=r_idx, column=c_idx).value = val
    for ci in range(1, (nk_ws.max_column or 19)+1):
        ws.column_dimensions[get_column_letter(ci)].width = 22

wb = Workbook()
last_row = build_main_sheet(wb.active)
build_nk_sheet(wb.create_sheet('NK_Konditionen'))
wb.save(OUT_XLSX)
print(f'\nGespeichert: {OUT_XLSX}')
print(f'  Sheet "GEZE GmbH":      {last_row} Zeilen')
print(f'  Sheet "NK_Konditionen": NK Geze.xlsx kopiert')
