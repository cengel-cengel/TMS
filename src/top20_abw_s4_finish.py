"""
Script 4: Add Methodik + VA-Mapping sheets, final save.
"""
import pickle, pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT  = Path('/home/user/TMS/output')
XLSX = OUT / 'top20_audit_abweichungsanalyse_2026.xlsx'

wb = load_workbook(str(XLSX))

BLUE_FILL  = PatternFill('solid', fgColor='1F4E79')
LBLUE_FILL = PatternFill('solid', fgColor='2E75B6')
LGREY_FILL = PatternFill('solid', fgColor='F2F2F2')

def hfont(bold=False, size=11, color='000000'):
    return Font(name='Calibri', bold=bold, size=size, color=color)

def thin():
    s = Side(style='thin')
    return Border(left=s, right=s, top=s, bottom=s)

# ═══════════════════════════════════════════════════════════════════════════════
# Sheet: Methodik
# ═══════════════════════════════════════════════════════════════════════════════
ws_m = wb.create_sheet('Methodik')
ws_m.column_dimensions['A'].width = 30
ws_m.column_dimensions['B'].width = 70

sections = [
    ('Analyseziel',
     'Vergleich der durchschnittlichen Erlöse pro Sendung vor und nach der TMS-Migration '
     '(Dinas → Microsoft AX, Go-Live 26.09.2025) für die 20 umsatzstärksten Kunden.'),
    ('Datenfilter',
     'Es werden ausschließlich Sendungen berücksichtigt, die eine Rechnungsnummer aufweisen '
     '(fakturierte Sendungen). Nicht-fakturierte Datensätze (Rechnungsnummer leer oder null) '
     'werden ausgeschlossen.'),
    ('Zeiträume',
     'PRE: alle Sendungen mit Leistungsdatum vor 26.09.2025.\n'
     'POST: alle Sendungen mit Leistungsdatum ab 26.09.2025.'),
    ('Verkehrsart-Normalisierung',
     'Da Dinas und AX unterschiedliche Verkehrsart-Bezeichnungen verwenden, wurden die '
     'Werte auf eine einheitliche Kategorie (va_norm) normiert. '
     'Siehe Sheet "VA Mapping" für Details.'),
    ('KPI: Durchschnittserlös',
     'Ø Erlös = Summe(Erloese) / Anzahl Sendungen im jeweiligen Zeitraum und Verkehrsart.\n'
     'Δ€ = Ø POST − Ø PRE.\n'
     'Δ% = Δ€ / |Ø PRE| × 100.'),
    ('KPI: DB II',
     'Deckungsbeitrag II = Erlöse − Kosten Gesamt (direkte Frachtkosten inkl. Diesel).\n'
     'Ø DB II wird analog zu Erlösen berechnet.'),
    ('Severity Score',
     'Severity = |Δ€| × max(n_pre, n_post).\n'
     'Gewichtet die absolute Erlösabweichung mit dem Sendungsvolumen, '
     'um die wirtschaftliche Relevanz zu priorisieren.'),
    ('Repräsentative Sendungen',
     'Je Kunde und Periode (PRE/POST) werden 5 Sendungen ausgewählt, '
     'deren Erlös am nächsten am Median-Erlös des Kunden liegt (Medianabstand).'),
    ('Farbkodierung',
     'Rot: Δ% ≤ −40% | Orange: −40% < Δ% ≤ −20% | Gelb: −20% < Δ% ≤ −5%\n'
     'Grün: Δ% ≥ +10% | Grau: n_post = 0 (Verkehr verloren)'),
    ('Erstellungsdatum', '02.04.2026'),
]

row = 1
ws_m.merge_cells('A1:B1')
c = ws_m.cell(1, 1, 'METHODIK — TOP-20 ABWEICHUNGSANALYSE')
c.font = hfont(bold=True, size=14, color='FFFFFF')
c.fill = BLUE_FILL
c.alignment = Alignment(horizontal='center')
ws_m.row_dimensions[1].height = 28

for i, (title, text) in enumerate(sections, 2):
    row = i + 1
    c_title = ws_m.cell(row, 1, title)
    c_title.font = hfont(bold=True)
    c_title.fill = LGREY_FILL
    c_title.border = thin()

    c_text = ws_m.cell(row, 2, text)
    c_text.alignment = Alignment(wrap_text=True, vertical='top')
    c_text.border = thin()
    ws_m.row_dimensions[row].height = max(30, text.count('\n') * 15 + 20)

# ═══════════════════════════════════════════════════════════════════════════════
# Sheet: VA Mapping
# ═══════════════════════════════════════════════════════════════════════════════
ws_va = wb.create_sheet('VA Mapping')
ws_va.column_dimensions['A'].width = 20
ws_va.column_dimensions['B'].width = 35
ws_va.column_dimensions['C'].width = 5
ws_va.column_dimensions['D'].width = 35

ws_va.merge_cells('A1:D1')
c = ws_va.cell(1, 1, 'VERKEHRSART-MAPPING: Dinas (PRE) → AX (POST)')
c.font = hfont(bold=True, size=13, color='FFFFFF')
c.fill = BLUE_FILL
c.alignment = Alignment(horizontal='center')
ws_va.row_dimensions[1].height = 26

headers = ['va_norm', 'Dinas (PRE) — Werte', '', 'AX (POST) — Werte']
for ci, h in enumerate(headers, 1):
    c = ws_va.cell(2, ci, h)
    c.font = hfont(bold=True, color='FFFFFF')
    c.fill = LBLUE_FILL
    c.border = thin()
    c.alignment = Alignment(horizontal='center')

mapping = [
    ('SA',
     "Verkehrsart='SA'",
     '→',
     "Verkehrsart='SA'"),
    ('Export',
     "Verkehrsart='Export', Detail='Stgt'\n+ Verkehrsart='Export', Detail='TP/DP'",
     '→',
     "Verkehrsart='Export', Detail='Stgt'\n+ Verkehrsart='Export', Detail='TP/DP'"),
    ('Charter',
     "Verkehrsart='Charter International', Detail='TP/DP'\n+ Verkehrsart='Charter International', Detail='Stgt'",
     '→',
     "Verkehrsart='Charter International', Detail='Charter '\n+ Verkehrsart='Charter International', Detail='TP/DP'"),
    ('Charter National',
     "Verkehrsart='Charter National'",
     '→',
     "Verkehrsart='Charter National'"),
    ('Belog SE',
     "Verkehrsart='Belog', Detail='Belog SE'",
     '→',
     "Verkehrsart='Belog', Detail='Belog'\n+ Verkehrsart='Belog', Detail='Belog SE'"),
    ('SE',
     "Verkehrsart='SE'",
     '→',
     "Verkehrsart='SE'"),
]

for ri, (norm, pre_val, arrow, post_val) in enumerate(mapping, 3):
    fill = PatternFill('solid', fgColor='F2F2F2') if ri % 2 == 1 else PatternFill(fill_type=None)
    for ci, val in enumerate([norm, pre_val, arrow, post_val], 1):
        c = ws_va.cell(ri, ci, val)
        c.fill = fill
        c.border = thin()
        c.alignment = Alignment(horizontal='center', wrap_text=True, vertical='top')
        if ci in (2, 4):
            ws_va.row_dimensions[ri].height = 35

print("Methodik and VA Mapping sheets added")
wb.save(str(XLSX))
print(f"Final saved: {XLSX}")

# ── Quick verification ────────────────────────────────────────────────────────
wb2 = load_workbook(str(XLSX), read_only=True)
print(f"Sheets ({len(wb2.sheetnames)}): {wb2.sheetnames}")
ov_ws = wb2['Abweichungsübersicht']
data_rows = [r for r in ov_ws.iter_rows(min_row=2, values_only=True) if any(c for c in r)]
print(f"Abweichungsübersicht data rows: {len(data_rows)}")
print("First row:", data_rows[0][:5] if data_rows else 'empty')
wb2.close()
