#!/usr/bin/env python3
"""
herma_cluster_step3_export.py
==============================
Schritt 3: Erstellt herma_dinas_vergleich.xlsx mit 4 Sheets:
  - Zusammenfassung
  - PRE Dinas Detail   (aus herma_vergleich_komplett.xlsx)
  - POST AX Detail     (aus herma_vergleich_komplett.xlsx)
  - Cluster-Vergleich  (neu, mit voller NK-Aufschlüsselung)

Eingabe:  output/herma_cluster_stats.csv
          output/herma_cluster_samples.csv
          output/herma_vergleich_komplett.xlsx
Ausgabe:  output/herma_dinas_vergleich.xlsx
"""

from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

KOMPLETT    = Path('/home/user/TMS/output/herma_vergleich_komplett.xlsx')
STATS_CSV   = Path('/home/user/TMS/output/herma_cluster_stats.csv')
SAMPLES_CSV = Path('/home/user/TMS/output/herma_cluster_samples.csv')
OUT_XLSX    = Path('/home/user/TMS/output/herma_dinas_vergleich.xlsx')

# ── Styles ────────────────────────────────────────────────────────────────────
def fill(hex_): return PatternFill('solid', fgColor=hex_)
def font(bold=False, color='000000', size=9, italic=False):
    return Font(bold=bold, color=color, size=size, italic=italic)

FILL_TITLE   = fill('1F497D')
FILL_CLUSTER = fill('2E75B6')
FILL_STATS   = fill('D6E4F7')
FILL_PRE_H   = fill('4472C4')
FILL_PRE     = fill('BDD7EE')
FILL_POST_H  = fill('C55A11')
FILL_POST    = fill('FCE4D6')
FILL_CTRL    = fill('E2EFDA')
FILL_ZS_H    = fill('2E4057')

THIN  = Side(border_style='thin',   color='AAAAAA')
MED   = Side(border_style='medium', color='555555')
BRD   = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BRD_H = Border(left=THIN, right=THIN, top=MED,  bottom=MED)
EUR   = '#,##0.00'

def wc(ws, row, col, value=None, f=None, fnt=None,
        al='left', fmt=None, brd=None):
    cell = ws.cell(row=row, column=col)
    if value is not None:
        cell.value = value
    if f:   cell.fill   = f
    if fnt: cell.font   = fnt
    if fmt: cell.number_format = fmt
    if brd: cell.border = brd
    cell.alignment = Alignment(horizontal=al, vertical='center')
    return cell

def merge(ws, row, c1, c2, val, f=None, fnt=None, al='left'):
    ws.merge_cells(start_row=row, start_column=c1,
                   end_row=row,   end_column=c2)
    wc(ws, row, c1, val, f, fnt, al)

# ── Spalten-Layout (PRE → POST Spalten-Mapping) ───────────────────────────────
# (header_label, pre_src, post_src, width)
# post_src=None → leer für POST-Zeilen
COLS = [
    ('RN',         'RN',          'RN',         10),
    ('Auftrag',    'Auftrag',     'Auftrag',     12),
    ('Datum',      'Leistungsdatum', None,        11),
    ('Land',       'Land',        'Land',          5),
    ('PLZ',        'PLZ',         'PLZ',           8),
    ('Zone',       'Zone',        'Zone',          8),
    ('Gew.band',   'Gew.band',    'Gew.band',     11),
    ('Ton. kg',    'Tonnage kg',  'Tonnage kg',    8),
    ('Soll €',     'SOLL EUR',    'SOLL EUR',      9),
    # NK — PRE: Dinas / POST: AX (gleiche Physik-Spalte, verschiedene Labels)
    ('Fracht €',   'Dinas Fracht','AX Fracht',    10),
    ('Diesel €',   'Diesel',      'Diesel',         9),
    ('Maut/SSD €', 'Maut/SSD',    'Maut',           9),
    ('Ausfuhr €',  'Ausfuhr',     'Neben',          9),
    ('Verzoll. €', 'Verzoll.',    None,             9),
    ('Zoll Bet. €','Zoll Bet.',   None,             9),
    ('Sulphur €',  'Sulphur',     None,             8),
    ('NK Pausch. €','NK Pausch',  'Versich.',      10),
    ('Sonstige €', 'Sonstige',    None,             9),
    ('GESAMT €',   'Dinas Gesamt','AX Gesamt',    11),
]
N = len(COLS)

PRE_HDR  = {i+1: f'D-{h}' if '€' in h else h
            for i, (h, _, _, _) in enumerate(COLS)}
PRE_HDR[10] = 'Dinas Fracht €'; PRE_HDR[11] = 'Dinas Diesel €'
PRE_HDR[12] = 'Dinas Maut €';   PRE_HDR[13] = 'Dinas Ausfuhr €'
PRE_HDR[14] = 'Dinas Verzoll. €'; PRE_HDR[15] = 'Dinas Zoll Bet. €'
PRE_HDR[16] = 'Dinas Sulphur €'; PRE_HDR[17] = 'Dinas NK Pausch. €'
PRE_HDR[18] = 'Dinas Sonstige €'; PRE_HDR[19] = 'Dinas GESAMT €'

POST_HDR = {i+1: h for i, (h, _, _, _) in enumerate(COLS)}
POST_HDR[10] = 'AX Fracht €'; POST_HDR[11] = 'AX Diesel €'
POST_HDR[12] = 'AX Maut €';   POST_HDR[13] = 'AX Nebengebühr €'
POST_HDR[14] = ''; POST_HDR[15] = ''
POST_HDR[16] = ''; POST_HDR[17] = 'AX Versich. €'
POST_HDR[18] = ''; POST_HDR[19] = 'AX GESAMT €'

def write_col_headers(ws, row, override, hdr_fill):
    for ci, (hdr, _, _, _) in enumerate(COLS, start=1):
        lbl = override.get(ci, hdr)
        wc(ws, row, ci, lbl, hdr_fill,
           font(bold=True, color='FFFFFF', size=8), 'center', brd=BRD_H)

def write_data_row(ws, row, src, src_idx, row_fill):
    for ci, col_def in enumerate(COLS, start=1):
        col_src = col_def[src_idx]  # 1=pre_src, 2=post_src
        val = src.get(col_src) if col_src else None
        if val is not None and pd.isna(val): val = None
        is_eur = '€' in col_def[0]
        wc(ws, row, ci, val, row_fill, font(size=9),
           'right' if is_eur or col_def[0] == 'Ton. kg' else 'left',
           EUR if is_eur else None, BRD)

# ── Sheet 1: Zusammenfassung ──────────────────────────────────────────────────
def write_zusammenfassung(ws, stats, pre, post):
    ws.title = 'Zusammenfassung'
    ws.column_dimensions['A'].width = 42
    ws.column_dimensions['B'].width = 22

    rows = [
        ('HERMA GmbH (423650) — Dinas PDF vs AX Vergleich', ''),
        ('', ''),
        ('── PRE (Dinas-System) ─────────────────────────', ''),
        ('BI-Zeilen PRE (mit Dinas-PDF gematcht)', len(pre)),
        ('Σ Dinas Fracht EUR', round(pd.to_numeric(pre['Dinas Fracht'], errors='coerce').sum(), 2)),
        ('Σ Dinas Gesamt EUR', round(pd.to_numeric(pre['Dinas Gesamt'], errors='coerce').sum(), 2)),
        ('Ø Dinas Gesamt EUR', round(pd.to_numeric(pre['Dinas Gesamt'], errors='coerce').mean(), 2)),
        ('', ''),
        ('── POST (AX-System) ────────────────────────────', ''),
        ('BI-Zeilen POST', len(post)),
        ('Σ AX Fracht EUR', round(pd.to_numeric(post['AX Fracht'], errors='coerce').sum(), 2)),
        ('Σ AX Gesamt EUR', round(pd.to_numeric(post['AX Gesamt'], errors='coerce').sum(), 2)),
        ('Ø AX Gesamt EUR', round(pd.to_numeric(post['AX Gesamt'], errors='coerce').mean(), 2)),
        ('', ''),
        ('── Cluster-Vergleich ────────────────────────────', ''),
        ('Unterfakturierungs-Cluster', len(stats)),
        ('Σ geschätzter Gesamtverlust EUR', round(stats['total_loss_est'].sum(), 0)),
        ('Schlimmster Cluster (Land)',
         stats.iloc[0]['Land'] + '-' + stats.iloc[0]['PLZ_pre']
         + ' ' + stats.iloc[0]['Gew_band'] if len(stats) else ''),
        ('', ''),
        ('── NK-Dateien ───────────────────────────────────', ''),
        ('Dinas PDF-Verzeichnis', 'data/extracted/v2/Herma/Rechnungen/Rechnungen DINAS/'),
        ('NK-Konditionen (docx)', '20251212_Herma_Nebenkosten_2026-2028.docx'),
        ('Hinweis NK docx', 'NK-Vergleich basiert auf tatsächl. PDF-abgerechneten NK'),
    ]

    for i, (k, v) in enumerate(rows, start=1):
        ws.cell(i, 1).value = k
        ws.cell(i, 2).value = v
        if v == '' and k.startswith('─'):
            ws.cell(i, 1).fill = fill('D6E4F7')
            ws.cell(i, 1).font = font(bold=True)
        elif i == 1:
            ws.cell(i, 1).font = font(bold=True, size=13)
        else:
            ws.cell(i, 2).alignment = Alignment(horizontal='right')

# ── Sheet 4: Cluster-Vergleich ─────────────────────────────────────────────────
def write_cluster_sheet(ws, stats, samples, top_n=20):
    ws.title = 'Cluster-Vergleich'
    merge(ws, 1, 1, N,
          'HERMA GmbH (423650) — Cluster-Vergleich: Unterfakturierung AX vs Dinas',
          FILL_TITLE, font(bold=True, color='FFFFFF', size=13), 'center')
    merge(ws, 2, 1, N,
          'Soll EUR aus Tarifmotor  |  PRE: Dinas NK  |  POST: AX NK  |  Grün = Kontroll-Paar',
          FILL_ZS_H, font(color='CCCCCC', size=9, italic=True), 'center')

    row = 2
    for _, cl in stats.head(top_n).iterrows():
        ckey = cl['_cluster']
        s    = samples[samples['_cluster'] == ckey]
        pre_s  = s[s['_system'] == 'PRE']
        post_s = s[s['_system'] == 'POST']
        kpre   = s[s['_system'] == 'KONTROLL_PRE']
        kpost  = s[s['_system'] == 'KONTROLL_POST']

        row += 1
        merge(ws, row, 1, N,
              f'  CLUSTER: Land={cl["Land"]}  |  PLZ-Prefix={cl["PLZ_pre"]}  '
              f'|  Gew.band={cl["Gew_band"]}  |  n PRE={int(cl["n_pre"])}  '
              f'|  n POST={int(cl["n_post"])}',
              FILL_CLUSTER, font(bold=True, color='FFFFFF', size=10))
        row += 1
        merge(ws, row, 1, N,
              f'  Ø Dinas: {cl["avg_dinas_ges"]:,.2f} €  |  Ø AX: {cl["avg_ax_ges"]:,.2f} €  '
              f'|  Δ: {cl["delta_ges"]:,.2f} € ({cl["delta_pct"]:+.1f}%)  '
              f'|  Gesch. Verlust: {cl["total_loss_est"]:,.0f} €  |  {cl["abw_grund"]}',
              FILL_STATS, font(size=9))

        # PRE
        row += 1
        merge(ws, row, 1, N, '  PRE — Dinas Rechnungen',
              FILL_PRE_H, font(bold=True, color='FFFFFF'))
        row += 1
        write_col_headers(ws, row, PRE_HDR, FILL_PRE_H)
        for _, pr in pre_s.iterrows():
            row += 1
            write_data_row(ws, row, pr, 1, FILL_PRE)

        # POST
        row += 1
        merge(ws, row, 1, N, '  POST — AX Rechnungen',
              FILL_POST_H, font(bold=True, color='FFFFFF'))
        row += 1
        write_col_headers(ws, row, POST_HDR, FILL_POST_H)
        for _, po in post_s.iterrows():
            row += 1
            write_data_row(ws, row, po, 2, FILL_POST)

        # Kontroll
        if not kpre.empty or not kpost.empty:
            row += 1
            merge(ws, row, 1, N,
                  '  ★ Kontroll-Sendung (bestes PRE/POST-Paar)',
                  FILL_CTRL, font(bold=True, color='1F5C00'))
            if not kpre.empty:
                row += 1
                wc(ws, row, 1, 'PRE', FILL_CTRL, font(bold=True, italic=True))
                pr = kpre.iloc[0]
                for ci, (h, psrc, _, _) in enumerate(COLS[1:], start=2):
                    v = pr.get(psrc) if psrc else None
                    if v is not None and pd.isna(v): v = None
                    is_eur = '€' in h
                    wc(ws, row, ci, v, FILL_CTRL, font(size=9),
                       'right' if is_eur else 'left',
                       EUR if is_eur else None, BRD)
            if not kpost.empty:
                row += 1
                wc(ws, row, 1, 'POST', FILL_CTRL, font(bold=True, italic=True))
                po = kpost.iloc[0]
                for ci, (h, _, postsrc, _) in enumerate(COLS[1:], start=2):
                    v = po.get(postsrc) if postsrc else None
                    if v is not None and pd.isna(v): v = None
                    is_eur = '€' in h
                    wc(ws, row, ci, v, FILL_CTRL, font(size=9),
                       'right' if is_eur else 'left',
                       EUR if is_eur else None, BRD)
        row += 1  # Leerzeile

    for ci, (_, _, _, w) in enumerate(COLS, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A3'
    return row

# ── Haupt ─────────────────────────────────────────────────────────────────────
def main():
    stats   = pd.read_csv(STATS_CSV,   sep=';')
    samples = pd.read_csv(SAMPLES_CSV, sep=';')
    pre     = pd.read_excel(KOMPLETT, sheet_name='PRE Dinas Detail', header=0)
    post    = pd.read_excel(KOMPLETT, sheet_name='POST AX Detail',   header=0)

    print(f'Cluster: {len(stats)}  |  Samples: {len(samples)}')

    wb = Workbook()

    # Sheet 1: Zusammenfassung
    ws_zs = wb.active
    write_zusammenfassung(ws_zs, stats, pre, post)

    # Sheet 2: PRE Dinas Detail (kopiert)
    ws_pre = wb.create_sheet('PRE Dinas Detail')
    ws_pre.append(list(pre.columns))
    for _, r in pre.iterrows():
        ws_pre.append(list(r))
    ws_pre.row_dimensions[1].height = 15
    for ci, _ in enumerate(pre.columns, start=1):
        ws_pre.column_dimensions[get_column_letter(ci)].width = 14

    # Sheet 3: POST AX Detail (kopiert)
    ws_post = wb.create_sheet('POST AX Detail')
    ws_post.append(list(post.columns))
    for _, r in post.iterrows():
        ws_post.append(list(r))
    for ci, _ in enumerate(post.columns, start=1):
        ws_post.column_dimensions[get_column_letter(ci)].width = 14

    # Sheet 4: Cluster-Vergleich
    ws_cv = wb.create_sheet('Cluster-Vergleich')
    last_row = write_cluster_sheet(ws_cv, stats, samples, top_n=20)

    wb.save(OUT_XLSX)
    print(f'Gespeichert: {OUT_XLSX}')
    print(f'  Sheet Cluster-Vergleich: {last_row} Zeilen')
    print(f'  Sheet PRE Dinas Detail:  {len(pre)} Zeilen')
    print(f'  Sheet POST AX Detail:    {len(post)} Zeilen')

if __name__ == '__main__':
    main()
